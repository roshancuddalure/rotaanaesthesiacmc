"""Check saved published workbook member rows against independent saved rota state.

Offline only: never connects to a database or changes source workbooks.
"""
from collections import Counter
from datetime import date
from pathlib import Path
import hashlib
import json
from openpyxl import load_workbook


def validate_member_cells(book, state):
    slots = {tuple(row['key']): row for row in state['slots']}
    assignments = [row for row in state['assignments'] if row['status'] in {'assigned', 'draft', 'confirmed'}]
    if not assignments:
        raise ValueError('A populated export fixture must have assignments')
    sheet = book['Final Rota']
    headers = [cell.value for cell in sheet[1]]
    required = ['Date', 'Unit', 'Duty', 'Slot Label', 'Assigned Member', 'Call Level',
                'Assignment Source', 'Override Reason', 'Template Status', 'Template Reason']
    if any(headers.count(name) != 1 for name in required):
        raise ValueError('Missing or duplicate final rota headers')
    indices = [headers.index(name) for name in required]
    actual = [tuple(row[i] for i in indices) for row in sheet.iter_rows(min_row=2, values_only=True)
              if any(value is not None for value in row)]
    expected = []
    counts = Counter()
    for assignment in assignments:
        key = tuple(assignment['slot'])
        slot = slots[key]
        expected.append((key[0], key[2], key[1], key[4], assignment['person'], key[3],
                         assignment['source'], assignment['override_reason'],
                         slot['template_status'], slot['template_reason']))
        weekend = date.fromisoformat(key[0]).weekday() >= 5
        for label, amount in [('Total Assignments', 1), ('24hr', int(slot['is_24hr'])),
                              ('Weekday Assignments', int(not weekend)), ('Weekend Assignments', int(weekend)),
                              ('Weekend 24hr', int(weekend and slot['is_24hr']))]:
            counts[assignment['person'], label] += amount
    if Counter(actual) != Counter(expected):
        raise ValueError('Final rota member/slot/source/status cells differ from saved assignments')
    duty = book['Duty Counts']
    labels = [cell.value for cell in duty[1]]
    observed = {}
    for row in duty.iter_rows(min_row=2, values_only=True):
        person = row[labels.index('Member')]
        if person in observed:
            raise ValueError('Duplicate member duty-count row')
        observed[person] = row
    for _, person in state['fixture_people']:
        if person not in observed:
            raise ValueError('Missing member duty-count row')
        for label in ['Total Assignments', '24hr', 'Weekday Assignments', 'Weekend Assignments', 'Weekend 24hr']:
            if observed[person][labels.index(label)] != counts[person, label]:
                raise ValueError('Duty count differs from independent saved assignments')
    return {'assigned_rows': len(expected), 'member_count_rows': len(observed),
            'exact_member_slot_source_status': True, 'independent_member_duty_counts': True}


def main():
    root = Path(__file__).resolve().parents[1]
    folder = root / 'docs/premium-ui/equivalence-publish'
    results = []
    rejected = 0
    for engine in ['chromium', 'firefox', 'webkit']:
        for width in [320, 1280]:
            for variant in ['ready', 'warnings']:
                for label in ['baseline', 'current']:
                    path = folder / f'{engine}-{width}-{variant}-{label}'
                    report = json.loads((path / 'report.json').read_text(encoding='utf8'))
                    workbook = path / 'final-rota.xlsx'
                    book = load_workbook(workbook)
                    try:
                        outcome = validate_member_cells(book, report['rota_before'])
                        # Deliberate changes must be rejected, even if old/new exports
                        # would otherwise share the same wrong values.
                        for sheet, cell, changed in [('Final Rota', 'G2', 'Wrong Member'),
                                                     ('Final Rota', 'A2', '2026-05-31'),
                                                     ('Final Rota', 'I2', 'wrong_source'),
                                                     ('Duty Counts', 'C2', 999)]:
                            original = book[sheet][cell].value
                            book[sheet][cell] = changed
                            try:
                                validate_member_cells(book, report['rota_before'])
                            except ValueError:
                                rejected += 1
                            else:
                                raise AssertionError(f'Negative control accepted: {sheet}!{cell}')
                            finally:
                                book[sheet][cell] = original
                    finally:
                        book.close()
                    results.append({'engine': engine, 'width': width, 'variant': variant, 'label': label,
                                    'workbook_sha256': hashlib.sha256(workbook.read_bytes()).hexdigest(),
                                    'browser_report_sha256': hashlib.sha256((path/'report.json').read_bytes()).hexdigest(),
                                    **outcome, 'passed': True})
    from datetime import datetime, timezone
    output = {'recorded_at_utc': datetime.now(timezone.utc).isoformat(),
              'scope': 'Offline independent populated published-workbook cell and count checks against retained browser-run saved state. Historical browser artifacts; no fresh publication run or production certification.',
              'whole_requirement_completion_proven': False, 'negative_controls_rejected': rejected,
              'results': results}
    (folder/'populated-member-cells-current.json').write_text(json.dumps(output, indent=2)+'\n', encoding='utf8')
    print(f'{len(results)} retained populated workbooks passed; {rejected} deliberate corruptions rejected')


if __name__ == '__main__':
    main()
