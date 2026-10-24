import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook
from frontend_publish_workbook_comparison import compare_publish_workbooks


class PublishWorkbookTests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.before = Path(temporary.name) / 'before'
        self.after = Path(temporary.name) / 'after'
        for path, instant in [(self.before, '2026-10-23T10:00:00'), (self.after, '2026-10-23T10:01:00')]:
            path.mkdir()
            (path / 'report.json').write_text(json.dumps({'approval_instants': [instant]}))
            book = Workbook()
            book.active.title = 'Summary'
            book.active['A9'] = 'Published at'
            book.active['B9'] = instant
            book.active['A10'] = 'Generated assignment'
            book.active['B10'] = 0
            book.save(path / 'final-rota.xlsx')
            book.close()

    def change(self, cell, value):
        path = self.after / 'final-rota.xlsx'
        book = load_workbook(path)
        book.active[cell] = value
        book.save(path)
        book.close()

    def test_distinct_exactly_validated_instants_allowed(self):
        self.assertTrue(compare_publish_workbooks(self.before, self.after)['artifacts_equal'])

    def test_wrong_recorded_instant_rejected(self):
        self.change('B9', '2026-10-23T10:02:00')
        with self.assertRaises(ValueError):
            compare_publish_workbooks(self.before, self.after)

    def test_changed_timestamp_label_rejected(self):
        self.change('A9', 'Different field')
        with self.assertRaises(ValueError):
            compare_publish_workbooks(self.before, self.after)

    def test_generated_text_never_hides_change(self):
        self.change('A10', 'Generated wrong assignment')
        self.assertFalse(compare_publish_workbooks(self.before, self.after)['artifacts_equal'])

    def test_zero_blank_difference_detected(self):
        self.change('B10', None)
        self.assertFalse(compare_publish_workbooks(self.before, self.after)['artifacts_equal'])

    def test_extra_column_detected(self):
        self.change('C10', 'extra')
        self.assertFalse(compare_publish_workbooks(self.before, self.after)['artifacts_equal'])


if __name__ == '__main__':
    unittest.main()
