"""Negative controls for the offline frontend artifact comparison gate."""
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook
from frontend_equivalence_artifacts import compare_artifacts


class ArtifactComparisonTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.before = Path(self.temp.name) / "before"
        self.after = Path(self.temp.name) / "after"
        for path in [self.before, self.after]:
            path.mkdir()
            (path / "report.json").write_text(json.dumps({
                "steps": ["saved"], "errors": [],
                "mutations": [{"method": "POST", "path": "/apply", "body": {"revision": 7}}]
            }), encoding="utf-8")
            book = Workbook()
            book.active.append(["Generated assignment", 0, "=B1+1"])
            book.save(path / "eagle-eye.xlsx")
            book.close()

    def change_book(self, mutate):
        path = self.after / "eagle-eye.xlsx"
        book = load_workbook(path)
        mutate(book)
        book.save(path)
        book.close()

    def change_report(self, mutate):
        path = self.after / "report.json"
        report = json.loads(path.read_text())
        mutate(report)
        path.write_text(json.dumps(report))

    def assert_different(self):
        self.assertFalse(compare_artifacts(self.before, self.after)["artifacts_equal"])

    def test_equal_is_bounded_not_certification(self):
        result = compare_artifacts(self.before, self.after)
        self.assertTrue(result["artifacts_equal"])
        self.assertFalse(result["whole_requirement_completion_proven"])

    def test_extra_column(self):
        self.change_book(lambda book: setattr(book.active["D1"], "value", "extra"))
        self.assert_different()

    def test_extra_row(self):
        self.change_book(lambda book: book.active.append(["extra"]))
        self.assert_different()

    def test_formula_change_even_without_cached_value(self):
        self.change_book(lambda book: setattr(book.active["C1"], "value", "=B1+2"))
        self.assert_different()

    def test_generated_word_does_not_hide_differences(self):
        self.change_book(lambda book: setattr(book.active["A1"], "value", "Generated wrong assignment"))
        self.assert_different()

    def test_zero_and_blank_differ(self):
        self.change_book(lambda book: setattr(book.active["B1"], "value", None))
        self.assert_different()

    def test_sheet_order(self):
        self.change_book(lambda book: book.create_sheet("extra", 0))
        self.assert_different()

    def test_revision_body(self):
        self.change_report(lambda report: report["mutations"][0]["body"].update(revision=8))
        self.assert_different()

    def test_error_is_not_success(self):
        self.change_report(lambda report: report["errors"].append("runtime failure"))
        self.assert_different()

    def test_empty_mutations_rejected(self):
        self.change_report(lambda report: report.update(mutations=[]))
        with self.assertRaises(ValueError):
            compare_artifacts(self.before, self.after)

    def test_missing_workbook_rejected(self):
        (self.after / "eagle-eye.xlsx").unlink()
        with self.assertRaises(OSError):
            compare_artifacts(self.before, self.after)


if __name__ == "__main__":
    unittest.main()
