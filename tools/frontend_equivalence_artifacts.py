"""Compare saved UI workflow artifacts without connecting to any database.

This checks recorded requests/outcomes and workbook values/formulas. It cannot
certify current source freshness, persisted database state, or unrecorded paths.
No timestamps, identifiers or cell text are silently excluded.
"""
import argparse
from datetime import date, datetime
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def fingerprint(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workbook_snapshot(path, *, cached=False):
    book = load_workbook(path, read_only=True, data_only=cached)
    try:
        sheets = []
        for sheet in book.worksheets:
            cells = []
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    if isinstance(value, (date, datetime)):
                        value = value.isoformat()
                    cells.append([cell.coordinate, cell.data_type, value])
            sheets.append({"name": sheet.title, "rows": sheet.max_row,
                           "columns": sheet.max_column, "cells": cells})
        return sheets
    finally:
        book.close()


def compare_artifacts(before, after):
    checks = []
    hashes = {}
    reports = []
    for label, directory in [("before", before), ("after", after)]:
        report_path = directory / "report.json"
        workbook_path = directory / "eagle-eye.xlsx"
        hashes[label] = {p.name: fingerprint(p) for p in [report_path, workbook_path]}
        report = json.loads(report_path.read_text(encoding="utf-8"))
        for key in ["steps", "errors", "mutations"]:
            if not isinstance(report.get(key), list):
                raise ValueError(f"{label} report needs a {key} list")
        if not report["steps"] or not report["mutations"]:
            raise ValueError(f"{label} report has no workflow outcomes or mutations")
        checks.append({"name": f"{label} successful recorded workflow",
                       "passed": not report["errors"] and not any(
                           str(step).startswith("FAILED:") for step in report["steps"])})
        reports.append(report)
    for key in ["steps", "mutations"]:
        checks.append({"name": f"recorded {key} exact equality",
                       "passed": reports[0][key] == reports[1][key]})
    for cached in [False, True]:
        checks.append({"name": "workbook cached values" if cached else "workbook values and formulas",
                       "passed": workbook_snapshot(before / "eagle-eye.xlsx", cached=cached)
                       == workbook_snapshot(after / "eagle-eye.xlsx", cached=cached)})
    return {"scope": "Saved artifacts only; no current-source or database equivalence certification.",
            "whole_requirement_completion_proven": False,
            "artifacts_equal": all(check["passed"] for check in checks),
            "checks": checks, "input_sha256": hashes}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("before", type=Path)
    parser.add_argument("after", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    try:
        result = compare_artifacts(args.before, args.after)
    except (OSError, ValueError, KeyError, TypeError) as error:
        parser.exit(2, f"Artifact comparison failed: {type(error).__name__}\n")
    output = json.dumps(result, indent=2) + "\n"
    if args.output:
        args.output.write_text(output, encoding="utf-8")
    print(output, end="")
    return 0 if result["artifacts_equal"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
