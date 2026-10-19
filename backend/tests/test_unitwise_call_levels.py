from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app import models  # noqa: F401
from app.db.session import Base
from app.models import Person
from app.services.members import StalePreviewError
from app.services.unitwise_call_levels import (
    extract_unitwise_call_levels,
    prefill_call_levels_from_unitwise,
    preview_call_level_prefill,
)


def make_unitwise(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append([None, "UNIT I", "UNIT II"])
    worksheet.append(["4th calls ", "Bernice", "Noble "])
    worksheet.append([None, "Amit Mathew ", None])
    worksheet.append(["Pain call", "May 1-15", "16-31"])
    worksheet.append(["Main campus", "Not A Call", None])
    workbook.save(path)


def test_extract_unitwise_call_levels_stops_at_non_call_sections(tmp_path: Path) -> None:
    path = tmp_path / "unitwise.xlsx"
    make_unitwise(path)

    entries = extract_unitwise_call_levels(path)

    assert [(entry.cleaned_name, entry.call_level) for entry in entries] == [
        ("Bernice", "4TH_CALL"),
        ("Noble", "4TH_CALL"),
        ("Amit Mathew", "4TH_CALL"),
    ]


def test_prefill_call_levels_previews_then_applies_only_certain_matches(tmp_path: Path) -> None:
    path = tmp_path / "unitwise.xlsx"
    make_unitwise(path)
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)

    with Session(engine) as session:
        session.add_all(
            [
                Person(canonical_name="Bernice Theodore Y"),
                Person(canonical_name="Noble E Cherian", call_level="3RD_CALL"),
                Person(canonical_name="Amit Mathew"),
            ]
        )
        session.commit()

        preview = preview_call_level_prefill(session, path)
        rows = {row["cleaned_name"]: row for row in preview.rows}
        # "Bernice" and "Noble" are only similar to a member: shown with candidates, never set.
        assert rows["Bernice"]["status"] == "needs_review"
        assert rows["Bernice"]["candidates"][0]["name"] == "Bernice Theodore Y"
        assert rows["Noble"]["action"] == "review"
        assert rows["Amit Mathew"]["action"] == "set"
        assert preview.changes == 1 and preview.mode == "preview"
        # The preview changes nothing.
        assert {p.canonical_name: p.call_level for p in session.scalars(select(Person))} == {
            "Bernice Theodore Y": None, "Noble E Cherian": "3RD_CALL", "Amit Mathew": None}

        with pytest.raises(StalePreviewError):
            prefill_call_levels_from_unitwise(session, path, token="stale")
        with pytest.raises(StalePreviewError):
            prefill_call_levels_from_unitwise(session, path)

        result = prefill_call_levels_from_unitwise(session, path, token=preview.preview_token)
        assert result.matched == 1 and result.mode == "applied"
        rows = {person.canonical_name: person.call_level for person in session.scalars(select(Person))}
    # Members without a confirmed row keep their level (nothing is cleared first).
    assert rows == {
        "Bernice Theodore Y": None,
        "Noble E Cherian": "3RD_CALL",
        "Amit Mathew": "4TH_CALL",
    }


def test_prefill_call_levels_does_not_cross_match_wrong_initial(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["DM/PDF", "Priyadharshini K"])
    path = tmp_path / "unitwise.xlsx"
    workbook.save(path)
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)

    with Session(engine) as session:
        session.add(Person(canonical_name="Priyadharshini S"))
        session.commit()

        preview = preview_call_level_prefill(session, path)
        result = prefill_call_levels_from_unitwise(session, path, token=preview.preview_token)

        assert result.matched == 0 and result.changes == 0
        person = session.scalars(select(Person)).one()
        assert person.call_level is None


def test_prefill_call_levels_flags_member_listed_at_two_levels(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["4th calls", "Amit Mathew"])
    worksheet.append(["5th calls", "Amit Mathew"])
    path = tmp_path / "unitwise.xlsx"
    workbook.save(path)
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)

    with Session(engine) as session:
        session.add(Person(canonical_name="Amit Mathew"))
        session.commit()

        preview = preview_call_level_prefill(session, path)
        assert {row["status"] for row in preview.rows} == {"conflict"}
        assert preview.changes == 0
