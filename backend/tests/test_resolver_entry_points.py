"""The golden set gives the same member decision through every import entry point.

tests/test_member_resolver.py proves the shared resolver against the golden set
(tests/fixtures/member_resolution_cases.csv). These tests drive the real preview builders of
each import with rows built from the same cases and check, per case, that

* the member the entry point puts on the row equals what ``MemberIndex.resolve`` gives for the
  same written name, with the context that entry point can carry (call level, date, duty
  area) and the purpose it uses ("live" or "history"); and
* no entry point auto-assigns a member outside the case's accepted set.

Context each entry point carries (from the code):
  review import units   call level (posting type) + row date            live
  review import leave   row date                                        live
  older leave import    row date                                        live
  older unit import     call level + row start date (+ duty area for special postings)  live
  past duties           call level (from duty) + duty date + duty area/campus           history
  member CSV            none                                            history
  member workbook       call level column, no date                      history
"""

import csv
import io
import json
from collections import defaultdict
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.db.session import Base
from app.models import Person, PersonAlias
from app.services import member_csv
from app.services.leave_import import preview_leave_import
from app.services.member_import import preview_member_import
from app.services.member_resolution_rules import seed_resolution_rules
from app.services.member_resolver import MemberIndex, ResolutionContext
from app.services.operational_import import normalize_rows
from app.services.past_duties import PAST_FIELDS, preview_past_duties
from app.services.unit_assignment_import import preview_unit_assignment_import

FIXTURES = Path(__file__).parent / "fixtures"
CASES = list(csv.DictReader(open(FIXTURES / "member_resolution_cases.csv", encoding="utf-8")))
LEVEL_CASES = [c for c in CASES if c["call_level"]]
NO_LEVEL_CASES = [c for c in CASES if not c["call_level"]]
DEFAULT_DAY = date(2026, 5, 15)  # for cases without a date (a completed past month)
PAST_DUTY = {"1ST_CALL": "MAIN_1ST_24HR", "2ND_CALL": "MAIN_2ND_24HR", "3RD_CALL": "MAIN_3RD_24HR",
             "4TH_CALL": "MAIN_4TH_24HR", "5TH_CALL": "FIFTH_CALL", "": "SCHELL_24HR"}
ORDINAL = {"1ST_CALL": "1st call", "2ND_CALL": "2nd call", "3RD_CALL": "3rd call",
           "4TH_CALL": "4th call", "5TH_CALL": "5th call"}


@pytest.fixture
def roster():
    engine = create_engine("sqlite+pysqlite:///:memory:", connect_args={"check_same_thread": False},
                           poolclass=StaticPool)
    Base.metadata.create_all(engine)
    db = Session(engine)
    snapshot = json.load(open(FIXTURES / "member_roster_snapshot.json", encoding="utf-8"))
    for p in snapshot["people"]:
        person = Person(canonical_name=p["name"], active_status=p["status"], call_level=p["call_level"])
        person.aliases = [PersonAlias(alias=a, source=s) for a, s in p["aliases"]]
        db.add(person)
    db.flush()
    seed_resolution_rules(db)
    db.commit()
    yield db
    db.close()
    engine.dispose()


def case_day(case) -> date:
    return date.fromisoformat(case["date"]) if case["date"] else DEFAULT_DAY


def by_month(cases):
    groups = defaultdict(list)
    for case in cases:
        groups[case_day(case).strftime("%Y-%m")].append(case)
    return sorted(groups.items())


def accepted(case) -> set[str]:
    return set(case["expected_member"].split(" / ")) if case["expected_member"] else set()


def single_name(case) -> bool:
    """Names that list several people are split into several rows by the older parsers."""
    return not any(sep in case["raw_name"] for sep in (",", "/", ";", " & ", " and "))


class Report:
    """Collects per-case disagreements so one failure lists every affected case."""

    def __init__(self, db, entry: str, purpose: str):
        self.index = MemberIndex.load(db)
        self.names = {str(p.id): p.canonical_name for p in self.index.people}
        self.entry, self.purpose = entry, purpose
        self.checked, self.differs, self.unsafe = 0, [], []

    def expected(self, raw: str, ctx: ResolutionContext) -> str | None:
        result = self.index.resolve(raw, ctx, purpose=self.purpose)
        return result.person_id if result.matched else None

    def check(self, case, got: str | None, ctx: ResolutionContext, *, compare: bool = True,
              name: str | None = None):
        """name: the text the entry point hands to the resolver when its parser cleans the cell
        (older leave/unit imports strip notes such as "(BP 16-30)"); defaults to the raw name."""
        self.checked += 1
        got = str(got) if got else None
        label = (case["raw_name"], case["call_level"], case["date"], case["unit"])
        if compare:
            want = self.expected(case["raw_name"] if name is None else name, ctx)
            if got != want:
                self.differs.append((*label, "entry=" + str(self.names.get(got)),
                                     "resolver=" + str(self.names.get(want))))
        if got and self.names[got] not in accepted(case):
            self.unsafe.append((*label, case["expected_member"] or case["expected_outcome"], self.names[got]))

    def assert_clean(self):
        assert self.checked, f"{self.entry}: no cases checked"
        assert self.unsafe == [], (
            f"{self.entry}: {len(self.unsafe)} automatic assignments outside the accepted set: {self.unsafe}")
        assert self.differs == [], (
            f"{self.entry}: {len(self.differs)} of {self.checked} cases differ from the resolver: {self.differs}")


# -- 1/2. review import (operational_import.normalize_rows) ------------------------------


def review_rows(cases, kind):
    rows = []
    for n, case in enumerate(cases):
        day = case_day(case).isoformat()
        values = {"member_name": case["raw_name"], "starts_on": day, "ends_on": day}
        if kind == "units":
            values["posting_type"] = case["call_level"]
            values["unit_name"] = case["unit"]
        rows.append({"key": f"CSV:{n + 2}", "source": {"sheet": "CSV", "row": n + 2}, "original": dict(values),
                     "values": values, "parse_issues": []})
    return rows


@pytest.mark.parametrize("kind,cases", [("units", LEVEL_CASES), ("leave", NO_LEVEL_CASES)])
def test_review_import(roster, kind, cases):
    report = Report(roster, f"review import {kind}", "live")
    for month, group in by_month(cases):
        result = normalize_rows(roster, kind, month, review_rows(group, kind), {"date_format": "auto"})
        assert len(result) == len(group)
        for case, row in zip(group, result):
            level = row["after"].get("posting_type") if kind == "units" else None
            assert kind == "leave" or level == case["call_level"], (case["raw_name"], level)
            report.check(case, row["after"]["person_id"], ResolutionContext.build(level, case_day(case)))
    report.assert_clean()


# -- 3. older leave import (leave_import.preview_leave_import) ---------------------------


def leave_csv(cases) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out, quoting=csv.QUOTE_ALL, lineterminator="\n")
    writer.writerow(["name", "start", "end"])
    for case in cases:
        writer.writerow([case["raw_name"], case_day(case).isoformat(), case_day(case).isoformat()])
    return out.getvalue().encode("utf-8")


def test_older_leave_import(roster):
    report = Report(roster, "older leave import", "live")
    for month, group in by_month(NO_LEVEL_CASES):
        preview = preview_leave_import(roster, "leave.csv", leave_csv(group), month)
        by_row = defaultdict(list)
        for row in preview["rows"]:
            by_row[row["row_number"]].append(row)
        for n, case in enumerate(group):
            rows = by_row.get(n + 2, [])
            if not rows:  # dropped as noise: nothing is assigned
                report.check(case, None, ResolutionContext.build(None, case_day(case)))
                continue
            for row in rows:
                report.check(case, row["person_id"], ResolutionContext.build(None, case_day(case)),
                             compare=len(rows) == 1 and single_name(case), name=row["cleaned_person_name"])
    report.assert_clean()


# -- 4. older unit import (unit_assignment_import.preview_unit_assignment_import) --------


def unit_text(cases) -> tuple[bytes, dict[int, object]]:
    lines, at_line = ["[UNIT: Test Unit]"], {}
    current = None
    for case in cases:
        if case["call_level"] != current:
            current = case["call_level"]
            lines.append(ORDINAL[current])
        lines.append(f"- {case['raw_name']}")
        at_line[len(lines)] = case
    return "\n".join(lines).encode("utf-8"), at_line


def test_older_unit_import(roster):
    report = Report(roster, "older unit import", "live")
    for month, group in by_month(LEVEL_CASES):
        group = sorted(group, key=lambda c: c["call_level"])
        content, at_line = unit_text(group)
        preview = preview_unit_assignment_import(roster, "units.txt", content, month)
        by_line = defaultdict(list)
        for row in preview["rows"]:
            by_line[row["row_number"]].append(row)
        for line, case in at_line.items():
            rows = by_line.get(line, [])
            if not rows:
                report.check(case, None, ResolutionContext.build(case["call_level"], date.fromisoformat(month + "-01")))
                continue
            for row in rows:
                assert row["posting_type"] == case["call_level"], (case["raw_name"], row["posting_type"])
                ctx = ResolutionContext.build(row["posting_type"], date.fromisoformat(row["starts_on"]))
                report.check(case, row["person_id"], ctx, compare=len(rows) == 1 and single_name(case),
                             name=row["cleaned_person_name"])
    report.assert_clean()


# -- 5. past duties (past_duties.preview_past_duties) ------------------------------------


def past_workbook(cases) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Past Duties"
    sheet.append(list(PAST_FIELDS))
    for case in cases:
        row = {"date": case_day(case).isoformat(), "name": case["raw_name"],
               "duty_type": PAST_DUTY[case["call_level"]], "campus": case["unit"]}
        sheet.append([row.get(field, "") for field in PAST_FIELDS])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def test_past_duties(roster):
    report = Report(roster, "past duties", "history")
    for month, group in by_month(CASES):
        preview = preview_past_duties(roster, "history.xlsx", past_workbook(group),
                                      {"month": month, "sheet": "Past Duties", "layout": "table"})
        assert len(preview["rows"]) == len(group)
        for case, row in zip(group, preview["rows"]):
            # Context as the import derives it: call level from the duty, area from the campus.
            ctx = ResolutionContext.build(case["call_level"] or None, case_day(case), case["unit"] or None)
            report.check(case, row["normalized"].get("person_id"), ctx)
    report.assert_clean()


# -- 6. member CSV (member_csv.build_preview) --------------------------------------------


def distinct(cases, key):
    seen, out = set(), []
    for case in cases:
        if key(case) not in seen:
            seen.add(key(case))
            out.append(case)
    return out


def test_member_csv(roster):
    cases = distinct(CASES, lambda c: c["raw_name"])
    out = io.StringIO()
    writer = csv.writer(out, quoting=csv.QUOTE_ALL, lineterminator="\n")
    writer.writerow(["canonical_name"])
    for case in cases:
        writer.writerow([case["raw_name"]])
    _, rows = member_csv.parse_file(out.getvalue().encode("utf-8"))
    preview = member_csv.build_preview(roster, rows, "update_existing")
    assert len(preview) == len(cases)
    report = Report(roster, "member CSV", "history")
    for case, row in zip(cases, preview):
        ctx = ResolutionContext.build()
        result = report.index.resolve(case["raw_name"], ctx, purpose="history")
        if not result.matched and any("confirm before assigning" in w for w in result.warnings):
            # Documented: an archived match is targeted so the checks can explain it (and
            # block reactivation); it is still the resolver's one proposed member.
            assert row["person_id"] == result.candidates[0].person_id, case["raw_name"]
            report.check(case, row["person_id"], ctx, compare=False)
            continue
        report.check(case, row["person_id"], ctx)
    report.assert_clean()


# -- 7. member workbook (member_import.preview_member_import) ----------------------------


def test_member_workbook(roster):
    cases = distinct(CASES, lambda c: (c["raw_name"], c["call_level"]))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Members"
    sheet.append(["name", "call_level"])
    for case in cases:
        sheet.append([case["raw_name"], ORDINAL.get(case["call_level"], "")])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    preview = preview_member_import(roster, "members.xlsx", buffer.getvalue(), {"sheet": "Members"})
    assert len(preview["rows"]) == len(cases)
    report = Report(roster, "member workbook", "history")
    for case, row in zip(cases, preview["rows"]):
        # A workbook row without an effective date describes the member as of today.
        ctx = ResolutionContext.build(case["call_level"] or None, date.today())
        dated = ResolutionContext.build(case["call_level"] or None, case_day(case))
        if case["date"] and report.expected(case["raw_name"], dated) != report.expected(case["raw_name"], ctx):
            # The case's own date gives a different answer than today; the workbook must
            # still agree with the resolver as of today.
            got = str(row["target_id"]) if row["target_id"] else None
            assert got == report.expected(case["raw_name"], ctx), (case["raw_name"], row["target_name"])
            continue
        report.check(case, row["target_id"], ctx)
    report.assert_clean()


# Dated rules apply to the workbook as of its effective date (else today): "Daniel" on 1st
# call is Joel Daniel from 2026-04-01.
def test_member_workbook_uses_current_dated_rules(roster):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Members"
    sheet.append(["name", "call_level"])
    sheet.append(["Daniel", "1st call"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    row = preview_member_import(roster, "members.xlsx", buffer.getvalue(), {"sheet": "Members"})["rows"][0]
    index = MemberIndex.load(roster)
    today = index.resolve("Daniel", ResolutionContext.build("1ST_CALL", date.today()), purpose="history")
    assert today.canonical_name == "Joel Daniel"
    assert row["target_name"] == today.canonical_name, row


# -- 1b. review import of an older unit-template file (read_rows -> legacy_rows) ---------


# The review import of an older unit-template file resolves the parser's cleaned name, so
# cells with notes ("Sanjana Annie Mathew (BP 16-30)") match as in the older unit import.
def test_review_import_of_unit_template_matches_older_unit_import(roster):
    from app.services.operational_import import read_rows

    differs, checked = [], 0
    for month, group in by_month(LEVEL_CASES):
        group = sorted(group, key=lambda c: c["call_level"])
        content, _ = unit_text(group)
        older = preview_unit_assignment_import(roster, "units.txt", content, month)["rows"]
        _, rows = read_rows(roster, "units", "units.txt", content, month, {})
        review = normalize_rows(roster, "units", month, rows, {"date_format": "auto"})
        assert len(review) == len(older)
        for old, new in zip(older, review):
            checked += 1
            if (old["person_id"] or None) != (new["after"]["person_id"] or None):
                differs.append((old["raw_person_name"], old["person_name"], new["person_match"]["canonical_name"]))
    assert checked and differs == [], f"{len(differs)} of {checked} rows differ: {differs}"
