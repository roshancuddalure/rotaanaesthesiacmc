import json
from datetime import date, datetime, timedelta
from io import BytesIO
from uuid import UUID
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.session import Base, get_db
from app.main import app
from app.models import (
    DutyAssignment,
    DutySlot,
    ImportBatch,
    ImportSourceRecord,
    LeaveRequest,
    Person,
    PersonAlias,
    PersonDesignation,
    PersonPosting,
    RotaPeriod,
    UserAccount,
)
from app.api.v1.auth import current_user, require_admin
from app.services.member_import import (
    FIELDS,
    apply_member_import,
    inspect_member_workbook,
    preview_member_import,
)


def workbook_bytes(rows, headers=FIELDS, sheet="Members", preamble=False):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    if preamble:
        worksheet.append(["Department Roster"])
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.fixture
def db():
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        yield session
    engine.dispose()


def apply_preview(db, content, options=None, filename="members.xlsx"):
    options = {"confirm_call_changes": True, **(options or {})}
    preview = preview_member_import(db, filename, content, options)
    result = apply_member_import(
        db, filename, content, options, preview["preview_hash"], "test-admin"
    )
    return preview, result


def test_csv_excel_reuploads_check_master_list_and_require_call_approval(db):
    original = workbook_bytes(
        [["Arjun Kumar", "", "first call"], ["Absent Doctor", "", "third call"]]
    )
    apply_preview(db, original)
    person = db.scalar(select(Person).where(Person.canonical_name == "Arjun Kumar"))
    original_id = person.id
    csv_data = b"\xef\xbb\xbfname,call_level\r\nDr. Arjun Kumar,second call\r\nNew Doctor,5\r\n"
    inspected = inspect_member_workbook("members.csv", csv_data)
    assert inspected["sheets"][0]["name"] == "CSV"
    assert inspected["sheets"][0]["mapping"] == {"name": 0, "call_level": 1}
    preview = preview_member_import(db, "members.csv", csv_data, {})
    assert preview["counts"]["update"] == 1
    assert preview["counts"]["create"] == 1
    assert preview["rows"][0]["changes"]["call_level"] == {"from": "1ST_CALL", "to": "2ND_CALL"}
    with pytest.raises(ValueError, match="Confirm"):
        apply_member_import(db, "members.csv", csv_data, {}, preview["preview_hash"], "admin")
    assert person.call_level == "1ST_CALL"
    _, result = apply_preview(db, csv_data, filename="members.csv")
    assert person.id == original_id and person.call_level == "2ND_CALL"
    assert (
        db.scalar(select(Person).where(Person.canonical_name == "Absent Doctor")).call_level
        == "3RD_CALL"
    )
    assert len(db.scalars(select(Person)).all()) == 3
    batch = db.get(ImportBatch, UUID(result["batch_id"]))
    assert batch.source_metadata["call_changes_confirmed"] is True
    assert batch.source_metadata["sheet"] == "CSV"
    preview, _ = apply_preview(db, csv_data, filename="members.csv")
    assert preview["counts"]["unchanged"] == 2
    assert len(db.scalars(select(Person)).all()) == 3
    blank = b"name,call_level\nArjun Kumar,\n"
    preview, _ = apply_preview(db, blank, filename="blank.csv")
    assert preview["counts"]["unchanged"] == 1
    assert person.call_level == "2ND_CALL"


def test_csv_matches_saved_aliases_and_checks_archived_members(db):
    person = Person(canonical_name="Arjun Kumar", active_status="archived")
    person.aliases.append(PersonAlias(alias="Arjun K"))
    db.add(person)
    db.commit()
    csv_data = b'name,call_level,aliases\nArjun K,4th call,"Arjun K;Dr. Arjun K"\n'
    # An archived member is never targeted silently; the reviewer must choose them.
    preview = preview_member_import(db, "members.csv", csv_data, {})
    assert preview["counts"]["needs_review"] == 1
    assert preview["rows"][0]["candidates"][0]["id"] == str(person.id)
    assert preview["rows"][0]["resolution_status"] == "needs_review"
    preview, _ = apply_preview(
        db, csv_data, {"resolutions": {"2": str(person.id)}}, filename="members.csv"
    )
    assert preview["counts"]["update"] == 1
    assert person.active_status == "archived"
    assert person.call_level == "4TH_CALL"
    assert len(db.scalars(select(Person)).all()) == 1
    dupes = b"name,call_level\nArjun Kumar,4th\nDr. Arjun Kumar,5th\n"
    preview = preview_member_import(
        db, "members.csv", dupes, {"resolutions": {"2": str(person.id)}}
    )
    assert preview["counts"]["needs_review"] == 1
    assert any("Duplicate name" in e for e in preview["rows"][1]["errors"])


@pytest.mark.parametrize(
    "content",
    [
        b"",
        b"\xff\xfehello",
        b"name,call_level\n\x00,first\n",
        b'name,call_level\n"unclosed,first\n',
    ],
)
def test_bad_csv_is_rejected(db, content):
    with pytest.raises(ValueError):
        inspect_member_workbook("members.csv", content)


def test_csv_semicolon_separator_quoted_aliases_and_fuzzy_review(db):
    person = Person(canonical_name="Arjun Kumar")
    db.add(person)
    db.commit()
    data = 'names;call levels;aliases\nArjun Kumaar;First;"Arjun K;Arjun Kumaar"\n'.encode()
    preview = preview_member_import(db, "members.csv", data, {})
    assert preview["counts"]["needs_review"] == 1
    assert preview["rows"][0]["candidates"][0]["id"] == str(person.id)
    preview, _ = apply_preview(
        db, data, {"resolutions": {"2": str(person.id)}}, filename="members.csv"
    )
    assert preview["counts"]["update"] == 1
    assert person.call_level == "1ST_CALL"
    assert "Arjun Kumaar" in {alias.alias for alias in person.aliases}


def test_normalized_creation_traceability_and_repeat_import(db):
    content = workbook_bytes(
        [
            [
                "Dr.  Arjun Kumar",
                "Asst. Professor",
                "first call",
                "ACTIVE",
                "A Kumar;Arjun K",
                date(2026, 10, 1),
            ]
        ]
    )
    preview = preview_member_import(db, "members.xlsx", content, {})
    assert preview["counts"]["create"] == 1
    assert db.scalars(select(Person)).all() == []
    assert db.scalars(select(ImportBatch)).all() == []
    _, result = apply_preview(db, content)
    person = db.scalars(select(Person)).one()
    assert person.canonical_name == "Arjun Kumar"
    assert person.call_level == "1ST_CALL"
    assert person.active_status == "active"
    assert person.designations[0].designation == "ASSISTANT PROFESSOR"
    assert person.designations[0].effective_from == date(2026, 10, 1)
    assert {a.alias for a in person.aliases} == {"A Kumar", "Arjun K", "Dr.  Arjun Kumar"}
    record = db.scalars(select(ImportSourceRecord)).one()
    assert record.row_index == 2 and record.sheet_name == "Members"
    assert record.normalized_id == person.id
    batch = db.get(ImportBatch, UUID(result["batch_id"]))
    assert batch.source_metadata["actor_id"] == "test-admin"
    second, _ = apply_preview(db, content)
    assert second["counts"]["unchanged"] == 1
    assert len(db.scalars(select(Person)).all()) == 1
    assert len(db.scalars(select(PersonDesignation)).all()) == 1
    assert len(db.scalars(select(PersonAlias)).all()) == 3


def test_alias_match_preserves_name_blank_fields_and_assignments(db):
    person = Person(
        canonical_name="Arjun Kumar",
        call_level="3RD_CALL",
        active_status="archived",
        archived_at=datetime(2026, 1, 1),
    )
    person.aliases.append(PersonAlias(alias="Arjun K"))
    period = RotaPeriod(name="2026-10", starts_on=date(2026, 10, 1), ends_on=date(2026, 10, 31))
    start = datetime(2026, 10, 10, 8)
    slot = DutySlot(
        rota_period=period,
        duty_date=start.date(),
        duty_type="MAIN_1ST_24HR",
        slot_label="Main",
        starts_at=start,
        ends_at=start + timedelta(days=1),
        is_24hr=True,
    )
    assignment = DutyAssignment(person=person, duty_slot=slot)
    leave = LeaveRequest(
        person=person, leave_type="annual", starts_on=date(2026, 10, 20), ends_on=date(2026, 10, 21)
    )
    posting = PersonPosting(person=person, posting_type="unit", starts_on=date(2026, 10, 1))
    db.add_all([person, assignment, leave, posting])
    db.commit()
    content = workbook_bytes([["Arjun K", "Professor", "", "", "", "2026-10-01"]])
    assert preview_member_import(db, "members.xlsx", content, {})["counts"]["needs_review"] == 1
    preview, _ = apply_preview(db, content, {"resolutions": {"2": str(person.id)}})
    assert preview["counts"]["update"] == 1
    assert person.canonical_name == "Arjun Kumar"
    assert person.call_level == "3RD_CALL"
    assert person.active_status == "archived"
    assert person.archived_at == datetime(2026, 1, 1)
    assert assignment.person_id == person.id
    assert db.scalars(select(DutyAssignment)).one().id == assignment.id
    assert db.scalars(select(LeaveRequest)).one().person_id == person.id
    assert db.scalars(select(PersonPosting)).one().person_id == person.id


def test_fuzzy_resolution_and_invalid_rows_must_be_resolved_or_skipped(db):
    person = Person(canonical_name="Arjun Kumar")
    db.add(person)
    db.commit()
    content = workbook_bytes(
        [["Arjun Kumaar", "", "second call"], ["DATE"], ["Other Doctor", "Professor", "nonsense"]]
    )
    preview = preview_member_import(db, "members.xlsx", content, {})
    assert preview["counts"]["needs_review"] == 3
    assert preview["rows"][0]["candidates"][0]["id"] == str(person.id)
    with pytest.raises(ValueError, match="Resolve or skip"):
        apply_member_import(db, "members.xlsx", content, {}, preview["preview_hash"], "admin")
    options = {"resolutions": {"2": str(person.id), "3": "skip", "4": "skip"}}
    preview, _ = apply_preview(db, content, options)
    assert preview["counts"]["skip"] == 2
    assert person.call_level == "2ND_CALL"
    assert "Arjun Kumaar" in {a.alias for a in person.aliases}
    assert len(db.scalars(select(Person)).all()) == 1
    assert (
        preview_member_import(
            db, "members.xlsx", content, {"resolutions": {"3": "skip", "4": "skip"}}
        )["counts"]["unchanged"]
        == 1
    )


def test_duplicates_and_cross_row_alias_conflicts(db):
    content = workbook_bytes([["Arjun Kumar", "", "", "", "Arjun K"], ["Arjun K"]])
    preview = preview_member_import(db, "members.xlsx", content, {})
    assert preview["rows"][1]["action"] == "needs_review"
    duplicate = workbook_bytes([["Arjun Kumar"], ["Dr. Arjun Kumar"]])
    preview = preview_member_import(db, "members.xlsx", duplicate, {})
    assert preview["counts"]["needs_review"] == 1
    # A skipped earlier row must not prevent importing the retained row.
    preview, _ = apply_preview(db, duplicate, {"resolutions": {"2": "skip"}})
    assert preview["counts"]["create"] == 1


def test_stale_preview_rejects_database_and_selection_changes(db):
    content = workbook_bytes([["Arjun Kumar"]])
    preview = preview_member_import(db, "members.xlsx", content, {})
    db.add(Person(canonical_name="Changed Member"))
    db.commit()
    with pytest.raises(ValueError, match="changed"):
        apply_member_import(db, "members.xlsx", content, {}, preview["preview_hash"], "admin")
    assert db.scalars(select(ImportBatch)).all() == []
    preview = preview_member_import(db, "members.xlsx", content, {})
    with pytest.raises(ValueError, match="changed"):
        apply_member_import(
            db,
            "members.xlsx",
            content,
            {"resolutions": {"2": "skip"}},
            preview["preview_hash"],
            "admin",
        )


@pytest.mark.parametrize(
    "row",
    [
        ["Valid Doctor", "Professor"],
        ["Valid Doctor", "Professor", "", "", "", "2026-11-01", "2026-10-01"],
        ["Valid Doctor", "Professor", "", "", "", "31/01/2026"],
        ["Valid Doctor", "", "1st call", "unknown"],
        ['=CONCAT("Some", "Doctor")'],
    ],
)
def test_invalid_dates_status_and_formulas_are_review_items(db, row):
    preview = preview_member_import(db, "members.xlsx", workbook_bytes([row]), {})
    assert preview["counts"]["needs_review"] == 1


def test_custom_headers_sheet_and_header_row(db):
    content = workbook_bytes(
        [["Valid Doctor", "fifth"]], headers=["Staff", "Tier"], sheet="Roster", preamble=True
    )
    inspection = inspect_member_workbook("members.xlsx", content, 2)
    assert inspection["sheets"][0]["headers"][:2] == ["Staff", "Tier"]
    options = {"header_row": 2, "sheet": "Roster", "mapping": {"name": 0, "call_level": 1}}
    preview, _ = apply_preview(db, content, options)
    assert preview["rows"][0]["row"] == 3
    assert db.scalars(select(Person)).one().call_level == "5TH_CALL"


def test_selected_worksheet_only_and_file_limits(db):
    workbook = Workbook()
    workbook.active.title = "Instructions"
    workbook.active.append(["Ignore this sheet"])
    sheet = workbook.create_sheet("Current Members")
    sheet.append(["Name"])
    sheet.append(["Valid Doctor"])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    content = buffer.getvalue()
    assert len(inspect_member_workbook("members.xlsx", content)["sheets"]) == 2
    preview, _ = apply_preview(db, content, {"sheet": "Current Members"})
    assert preview["counts"]["create"] == 1
    with pytest.raises(ValueError, match="5 MB"):
        inspect_member_workbook("members.xlsx", b"x" * (5 * 1024 * 1024 + 1))
    with pytest.raises(ValueError, match="5000"):
        preview_member_import(db, "members.xlsx", workbook_bytes([["Valid Doctor"]] * 5001), {})


def test_import_atomic_rollback_on_write_failure(db, monkeypatch):
    content = workbook_bytes([["First Doctor"], ["Second Doctor"]])
    preview = preview_member_import(db, "members.xlsx", content, {})

    def fail_commit():
        raise RuntimeError("Simulated write failure")

    monkeypatch.setattr(db, "commit", fail_commit)
    with pytest.raises(RuntimeError):
        apply_member_import(db, "members.xlsx", content, {}, preview["preview_hash"], "admin")
    assert db.scalars(select(Person)).all() == []
    assert db.scalars(select(ImportBatch)).all() == []


def test_api_permissions_template_and_full_workflow(db):
    factory = sessionmaker(bind=db.get_bind())
    db.add(
        UserAccount(
            id=UUID(int=1),
            username="test-admin",
            display_name="Department Admin",
            role="superadmin",
            password_hash="test-only",
        )
    )
    db.commit()

    def override_db():
        with factory() as session:
            yield session

    app.dependency_overrides[get_db] = override_db
    client = TestClient(app)
    prefix = "/api/v1/admin/members/import"
    try:
        assert client.get(prefix + "/template").status_code == 401
        assert client.post(prefix + "/preview").status_code == 401
        assert client.get(prefix + "/history").status_code == 401
        app.dependency_overrides[current_user] = lambda: UserAccount(role="rota_board_member")
        assert client.get(prefix + "/template").status_code == 403
        assert client.get(prefix + "/history").status_code == 403
        del app.dependency_overrides[current_user]
        app.dependency_overrides[require_admin] = lambda: UserAccount(
            id=UUID(int=1), role="superadmin"
        )
        template = client.get(prefix + "/template")
        assert template.status_code == 200
        workbook = load_workbook(BytesIO(template.content))
        assert list(next(workbook.active.values)) == list(FIELDS)
        workbook.close()
        content = workbook_bytes([["Valid Doctor"]])
        files = {"file": ("members.xlsx", content)}
        inspect = client.post(prefix + "/inspect", files=files)
        assert inspect.status_code == 200
        preview = client.post(prefix + "/preview", files=files)
        assert preview.status_code == 200
        applied = client.post(
            prefix + "/apply",
            files=files,
            data={"options": "{}", "preview_hash": preview.json()["preview_hash"]},
        )
        assert applied.status_code == 200, applied.text
        assert applied.json()["counts"]["create"] == 1
        csv_template = client.get(prefix + "/template?format=csv")
        assert csv_template.status_code == 200
        assert csv_template.content.decode("utf-8-sig").strip() == "name,call_level"
        history = client.get(prefix + "/history")
        assert len(history.json()) == 1
        assert history.json()[0]["filename"] == "members.xlsx"
        assert history.json()[0]["actor"] == "Department Admin"
        detail = client.get(prefix + "/history/" + applied.json()["batch_id"])
        assert detail.status_code == 200
        assert detail.json()["rows"][0]["raw"]["name"] == "Valid Doctor"
        assert detail.json()["rows"][0]["action"] == "create"
        unrelated = ImportBatch(source_filename="rota.xlsx", import_kind="historical")
        db.add(unrelated)
        db.commit()
        assert client.get(prefix + "/history/" + str(unrelated.id)).status_code == 404
        assert len(client.get(prefix + "/history").json()) == 1
        csv_file = {"file": ("calls.csv", b"name,call_level\nValid Doctor,second call\n")}
        csv_preview = client.post(prefix + "/preview", files=csv_file)
        assert csv_preview.status_code == 200
        assert csv_preview.json()["counts"]["update"] == 1
        payload = {"options": "{}", "preview_hash": csv_preview.json()["preview_hash"]}
        assert client.post(prefix + "/apply", files=csv_file, data=payload).status_code == 400
        payload["options"] = json.dumps({"confirm_call_changes": True})
        applied_csv = client.post(prefix + "/apply", files=csv_file, data=payload)
        assert applied_csv.status_code == 200, applied_csv.text
        csv_history = client.get(prefix + "/history/" + applied_csv.json()["batch_id"]).json()
        assert csv_history["batch"]["call_changes_confirmed"] is True
        assert csv_history["batch"]["sheet"] == "CSV"
        assert csv_history["rows"][0]["changes"]["call_level"] == {"from": None, "to": "2ND_CALL"}
        unchanged_csv = client.post(prefix + "/preview", files=csv_file)
        assert unchanged_csv.json()["counts"]["unchanged"] == 1
        broken = BytesIO()
        with ZipFile(BytesIO(content)) as source, ZipFile(broken, "w") as target:
            for item in source.infolist():
                target.writestr(
                    item.filename,
                    b"<broken"
                    if item.filename == "xl/worksheets/sheet1.xml"
                    else source.read(item.filename),
                )
        assert (
            client.post(
                prefix + "/inspect", files={"file": ("bad.xlsx", broken.getvalue())}
            ).status_code
            == 400
        )
        assert (
            client.post(prefix + "/preview", files=files, data={"options": "[]"}).status_code == 400
        )
        assert (
            client.post(prefix + "/preview", files={"file": ("bad.xlsx", b"bad")}).status_code
            == 400
        )
        assert (
            client.post(prefix + "/preview", files={"file": ("bad.xls", b"bad")}).status_code == 400
        )
        assert (
            client.post(
                prefix + "/preview",
                files=files,
                data={"options": json.dumps({"mapping": {"name": -1}})},
            ).status_code
            == 400
        )
    finally:
        app.dependency_overrides.clear()
