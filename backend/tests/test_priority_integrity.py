"""Regression cases for the audit's initial permission and data-integrity fixes."""
from datetime import date
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.db.session import Base, get_db
from app.main import app
from app.models import ImportBatch, ImportSourceRecord, LeaveRequest, Person, PersonPosting, Unit, UserAccount
from app.services.auth import hash_secret
from app.services.unit_assignment_import import clear_unit_board_month, apply_unit_assignment_import
from app.services.rota_setup import monthly_setup, update_scope_units
from app.services.member_import import preview_member_import, apply_member_import
from app.services.past_duties import preview_past_duties, apply_past_duties
from app.models import DutyAssignment


@pytest.fixture
def secured_client():
    engine = create_engine('sqlite+pysqlite:///:memory:', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(engine)
    with factory() as db:
        for role in ['superadmin','computer_admin','hod','rota_board_member']:
            db.add(UserAccount(username=role, display_name=role, role=role, password_hash=hash_secret('password123')))
        db.commit()
    def database():
        with factory() as db:
            yield db
    previous = dict(app.dependency_overrides)
    app.dependency_overrides.clear()
    app.dependency_overrides[get_db] = database
    try:
        yield TestClient(app), factory
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(previous)
        engine.dispose()


def login(client, role):
    result = client.post('/api/v1/auth/sign-in', json={'username':role,'password':'password123'})
    assert result.status_code == 200
    return result.json()


@pytest.mark.parametrize('method,path', [
    ('GET','/api/v1/admin/members'), ('POST','/api/v1/admin/members'),
    ('POST','/api/v1/admin/members/cleanup-invalid'), ('POST','/api/v1/admin/members/reset-from-trusted-roster'),
    ('GET','/api/v1/admin/mappings'), ('POST','/api/v1/admin/mappings/scan-historical'),
    ('GET','/api/v1/analysis/dashboard'), ('GET','/api/v1/analysis/preflight'),
    ('GET','/api/v1/analysis/manual-review'), ('POST','/api/v1/admin/imports/historical/run'),
])
def test_legacy_endpoints_reject_anonymous(secured_client, method, path):
    client,_ = secured_client
    assert client.request(method,path,json={}).status_code == 401


@pytest.mark.parametrize('role',['hod','rota_board_member'])
def test_non_admin_can_read_members_but_cannot_mutate(secured_client, role):
    client,_ = secured_client
    token=login(client,role)['token']; headers={'Authorization':f'Bearer {token}'}
    assert client.get('/api/v1/admin/members',headers=headers).status_code == 200
    assert client.post('/api/v1/admin/members',headers=headers,json={'canonical_name':'Valid Doctor'}).status_code == 403
    assert client.post('/api/v1/admin/members/cleanup-invalid',headers=headers).status_code == 403
    assert client.get('/api/v1/admin/mappings',headers=headers).status_code == 403


def test_bootstrap_removed_and_legacy_rebuild_retired(secured_client):
    client,_ = secured_client
    assert client.post('/api/v1/auth/seed-superadmin').status_code in {404,405}
    token=login(client,'superadmin')['token']
    assert client.post('/api/v1/admin/imports/historical/run',headers={'Authorization':f'Bearer {token}'}).status_code == 410


def test_recovery_is_private_and_revokes_sessions(secured_client):
    client,_ = secured_client
    known=client.post('/api/v1/auth/forgot-password',json={'username':'hod'})
    unknown=client.post('/api/v1/auth/forgot-password',json={'username':'does-not-exist'})
    assert known.json() == unknown.json()
    assert 'reset_token' not in known.json()
    hod=login(client,'hod'); admin=login(client,'superadmin')
    path=f"/api/v1/auth/accounts/{hod['user']['id']}/recovery-code"
    assert client.post(path).status_code == 401
    assert client.post(path,headers={'Authorization':f"Bearer {hod['token']}"}).status_code == 403
    headers={'Authorization':f"Bearer {admin['token']}"}
    first=client.post(path,headers=headers).json()['reset_token']
    second=client.post(path,headers=headers).json()['reset_token']
    assert client.post('/api/v1/auth/reset-password',json={'token':first,'new_password':'newpassword123'}).status_code == 400
    assert client.post('/api/v1/auth/reset-password',json={'token':second,'new_password':'short'}).status_code == 400
    assert client.post('/api/v1/auth/reset-password',json={'token':second,'new_password':'newpassword123'}).status_code == 200
    assert client.get('/api/v1/auth/me',headers={'Authorization':f"Bearer {hod['token']}"}).status_code == 401
    assert client.post('/api/v1/auth/reset-password',json={'token':second,'new_password':'otherpassword'}).status_code == 400
    assert client.post('/api/v1/auth/sign-in',json={'username':'hod','password':'newpassword123'}).status_code == 200


def test_computer_admin_cannot_create_or_recover_superadmin(secured_client):
    client,_=secured_client
    admin=login(client,'computer_admin'); chief=login(client,'superadmin')
    headers={'Authorization':f"Bearer {admin['token']}"}
    payload={'username':'new-chief','display_name':'Chief','role':'superadmin','password':'password123'}
    assert client.post('/api/v1/auth/accounts',headers=headers,json=payload).status_code == 403
    assert client.post(f"/api/v1/auth/accounts/{chief['user']['id']}/recovery-code",headers=headers).status_code == 403
    payload['role']='hod';payload['password']='short'
    assert client.post('/api/v1/auth/accounts',headers=headers,json=payload).status_code == 400


@pytest.mark.parametrize('payload', [
    {'canonical_name':''}, {'canonical_name':'Member 123'},
    {'canonical_name':'Valid Doctor','active_status':'unknown'},
    {'canonical_name':'Valid Doctor','call_level':'SICU'},
])
def test_manual_member_validation_matches_import_rules(secured_client,payload):
    client,_=secured_client
    token=login(client,'superadmin')['token']
    assert client.post('/api/v1/admin/members',headers={'Authorization':f'Bearer {token}'},json=payload).status_code == 422


def test_duplicate_create_preserves_existing_member(secured_client):
    client,_=secured_client
    token=login(client,'superadmin')['token'];headers={'Authorization':f'Bearer {token}'}
    original=client.post('/api/v1/admin/members',headers=headers,json={'canonical_name':'Dr. Valid Doctor','active_status':'historical','call_level':'first call'}).json()
    assert original['canonical_name']=='Valid Doctor'
    duplicate=client.post('/api/v1/admin/members',headers=headers,json={'canonical_name':'valid doctor','active_status':'active','call_level':'2ND_CALL'})
    assert duplicate.status_code==409
    saved=next(m for m in client.get('/api/v1/admin/members',headers=headers).json() if m['id']==original['id'])
    assert saved['active_status']=='historical' and saved['call_level']=='1ST_CALL'


def test_month_replacement_preserves_both_sides_and_other_sources(secured_client):
    _,factory=secured_client
    with factory() as db:
        person=Person(canonical_name='Valid Doctor');unit=Unit(code='UNIT_I',name='Unit I')
        db.add_all([person,unit]);db.flush()
        for source,start,end in [('unit_board',date(2026,1,1),None),('unit_board',date(2026,10,3),date(2026,10,5)),('manual',date(2026,1,1),None)]:
            db.add(PersonPosting(person_id=person.id,unit_id=unit.id,posting_type='1ST_CALL',starts_on=start,ends_on=end,source=source))
        db.commit()
        assert clear_unit_board_month(db,'2026-10')==2
        db.commit()
        rows=list(db.scalars(select(PersonPosting)))
        periods={(p.source,p.starts_on,p.ends_on) for p in rows}
        assert periods=={('unit_board',date(2026,1,1),date(2026,9,30)),('unit_board',date(2026,11,1),None),('manual',date(2026,1,1),None)}


def test_invalid_replacement_keeps_postings(secured_client):
    _,factory=secured_client
    with factory() as db:
        p=Person(canonical_name='Valid Doctor');db.add(p);db.flush()
        db.add(PersonPosting(person_id=p.id,posting_type='1ST_CALL',starts_on=date(2026,1,1),source='unit_board'));db.commit()
        w=Workbook();w.active.append(['Posting','Unit I']);w.active.append(['first call','Unknown Person']);b=BytesIO();w.save(b)
        with pytest.raises(ValueError,match='at least one valid'):
            apply_unit_assignment_import(db,'bad.xlsx',b.getvalue(),'2026-10',replace_existing=True)
        assert len(list(db.scalars(select(PersonPosting))))==1


@pytest.mark.parametrize('start,end,expected', [
    (date(2026,9,2),date(2026,10,12),[(date(2026,9,2),date(2026,9,30))]),
    (date(2026,10,12),date(2026,11,2),[(date(2026,11,1),date(2026,11,2))]),
    (date(2026,10,1),date(2026,10,31),[]),
    (date(2026,11,1),None,[(date(2026,11,1),None)]),
    (date(2026,9,1),date(2026,9,30),[(date(2026,9,1),date(2026,9,30))]),
])
def test_replacement_date_boundaries(secured_client,start,end,expected):
    _,factory=secured_client
    with factory() as db:
        p=Person(canonical_name='Boundary Doctor');db.add(p);db.flush()
        db.add(PersonPosting(person_id=p.id,posting_type='1ST_CALL',starts_on=start,ends_on=end,source='unit_board'));db.commit()
        clear_unit_board_month(db,'2026-10');db.commit()
        assert [(p.starts_on,p.ends_on) for p in db.scalars(select(PersonPosting))]==expected


def test_locked_scope_cannot_change_without_unlock(secured_client):
    _,factory=secured_client
    with factory() as db:
        one=Unit(code='ONE',name='One');two=Unit(code='TWO',name='Two');db.add_all([one,two]);db.commit()
        _,scope=monthly_setup(db,'2026-10')
        update_scope_units(db,scope,[one.id],[two.id],True,True,'Approved')
        with pytest.raises(ValueError,match='Unlock'):
            update_scope_units(db,scope,[two.id],[one.id],True,True,'Still locked')
        with pytest.raises(ValueError,match='reason'):
            update_scope_units(db,scope,[two.id],[one.id],True,False,'')
        saved=update_scope_units(db,scope,[one.id],[two.id],True,True,'Approved')
        assert saved.is_locked
        saved=update_scope_units(db,scope,[two.id],[one.id],True,False,'Board authorized change')
        assert not saved.is_locked


@pytest.mark.parametrize('path',[
    '/api/v1/leave/import-preview?month=2026-10',
    '/api/v1/leave/import-apply?month=2026-10',
    '/api/v1/unit-management/import-preview?month=2026-10',
    '/api/v1/unit-management/import-apply?month=2026-10',
])
def test_large_upload_rejected_before_parsing(secured_client,path):
    client,_=secured_client
    token=login(client,'superadmin')['token']
    response=client.post(path,headers={'Authorization':f'Bearer {token}'},files={'file':('large.xlsx',b'x'*(5*1024*1024+1),'application/octet-stream')})
    assert response.status_code==400
    assert '5 MB' in response.json()['detail']


def test_leave_apply_requires_current_reviewed_source_and_master_state(secured_client):
    client,factory=secured_client
    token=login(client,'superadmin')['token']; headers={'Authorization':f'Bearer {token}'}
    with factory() as db:
        db.add(Person(canonical_name='Import Doctor',call_level='1ST_CALL'));db.commit()
    source=b'name,starts_on,ends_on,leave_type\nImport Doctor,2026-10-20,2026-10-21,ANNUAL_LEAVE\n'
    files={'file':('leave.csv',source,'text/csv')}
    preview=client.post('/api/v1/leave/import-preview?month=2026-10',headers=headers,files=files)
    assert preview.status_code==200
    reviewed=preview.json()['preview_hash']
    path='/api/v1/leave/import-apply?month=2026-10'
    assert client.post(path,headers=headers,files=files).status_code==400
    changed={'file':('leave.csv',source.replace(b'10-20',b'10-22'),'text/csv')}
    assert client.post(path,headers=headers,files=changed,data={'preview_hash':reviewed}).status_code==409
    with factory() as db:
        person=db.scalar(select(Person).where(Person.canonical_name=='Import Doctor'))
        person.call_level='2ND_CALL';db.commit()
    assert client.post(path,headers=headers,files=files,data={'preview_hash':reviewed}).status_code==409
    current=client.post('/api/v1/leave/import-preview?month=2026-10',headers=headers,files=files).json()['preview_hash']
    applied=client.post(path,headers=headers,files=files,data={'preview_hash':current})
    assert applied.status_code==200 and not applied.json()['replayed']
    replay=client.post(path,headers=headers,files=files,data={'preview_hash':current})
    assert replay.status_code==200 and replay.json()['replayed']
    assert replay.json()['batch_id']==applied.json()['batch_id']
    with factory() as db:
        assert len(list(db.scalars(select(LeaveRequest))))==1
        assert len(list(db.scalars(select(ImportBatch))))==1
        records=list(db.scalars(select(ImportSourceRecord)))
        assert len(records)==1 and records[0].normalized_id is not None


def test_unit_apply_rejects_options_changed_after_review(secured_client):
    client,factory=secured_client
    token=login(client,'superadmin')['token'];headers={'Authorization':f'Bearer {token}'}
    with factory() as db:
        db.add_all([Person(canonical_name='Import Doctor',call_level='1ST_CALL'),Unit(code='UNIT_I',name='Unit I')]);db.commit()
    w=Workbook();w.active.append(['Posting','Unit I']);w.active.append(['1st call','Import Doctor']);b=BytesIO();w.save(b)
    files={'file':('units.xlsx',b.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    preview=client.post('/api/v1/unit-management/import-preview?month=2026-10',headers=headers,files=files)
    assert preview.status_code==200
    reviewed=preview.json()['preview_hash']
    assert client.post('/api/v1/unit-management/import-apply?month=2026-10&replace_existing=true',headers=headers,files=files,data={'preview_hash':reviewed}).status_code==409
    applied=client.post('/api/v1/unit-management/import-apply?month=2026-10',headers=headers,files=files,data={'preview_hash':reviewed})
    assert applied.status_code==200 and applied.json()['created_rows']==1
    replay=client.post('/api/v1/unit-management/import-apply?month=2026-10',headers=headers,files=files,data={'preview_hash':reviewed})
    assert replay.status_code==200 and replay.json()['replayed']
    assert replay.json()['batch_id']==applied.json()['batch_id']
    with factory() as db:
        assert len(list(db.scalars(select(PersonPosting))))==1
        assert len(list(db.scalars(select(ImportBatch))))==1


def test_leave_and_audit_commit_are_atomic(secured_client,monkeypatch):
    client,factory=secured_client
    token=login(client,'superadmin')['token']; headers={'Authorization':f'Bearer {token}'}
    with factory() as db:
        db.add(Person(canonical_name='Atomic Doctor',call_level='1ST_CALL'));db.commit()
    files={'file':('leave.csv',b'name,starts_on,ends_on\nAtomic Doctor,2026-10-20,2026-10-21\n','text/csv')}
    reviewed=client.post('/api/v1/leave/import-preview?month=2026-10',headers=headers,files=files).json()['preview_hash']
    def failure(*args,**kwargs):
        raise RuntimeError('audit failure')
    monkeypatch.setattr('app.api.v1.leave.record_workflow_batch',failure)
    with pytest.raises(RuntimeError,match='audit failure'):
        client.post('/api/v1/leave/import-apply?month=2026-10',headers=headers,files=files,data={'preview_hash':reviewed})
    with factory() as db:
        assert list(db.scalars(select(LeaveRequest)))==[]
        assert list(db.scalars(select(ImportBatch)))==[]


def test_member_and_past_duty_retries_return_one_batch_each(secured_client):
    _,factory=secured_client
    with factory() as db:
        data=b'name,call_level\nReplay Doctor,first call\n'
        preview=preview_member_import(db,'members.csv',data,{})
        first=apply_member_import(db,'members.csv',data,{},preview['preview_hash'],'internal-test')
        replay=apply_member_import(db,'members.csv',data,{},preview['preview_hash'],'internal-test')
        assert replay['replayed'] and replay['batch_id']==first['batch_id']
        assert len(list(db.scalars(select(Person))))==1
        w=Workbook();w.active.title='Past Duties';w.active.append(['date','name','duty_type']);w.active.append(['2026-09-20','Replay Doctor','MAIN_1ST_24HR']);b=BytesIO();w.save(b)
        options={'month':'2026-09','sheet':'Past Duties','layout':'table'}
        preview=preview_past_duties(db,'history.xlsx',b.getvalue(),options)
        first=apply_past_duties(db,'history.xlsx',b.getvalue(),options,preview['preview_hash'],'internal-test')
        replay=apply_past_duties(db,'history.xlsx',b.getvalue(),options,preview['preview_hash'],'internal-test')
        assert replay['replayed'] and replay['batch_id']==first['batch_id']
        assert len(list(db.scalars(select(DutyAssignment))))==1
        assert len(list(db.scalars(select(ImportBatch))))==2


def test_distinct_member_files_can_have_independent_unchanged_batches(secured_client):
    _,factory=secured_client
    with factory() as db:
        db.add(Person(canonical_name='Existing Doctor',call_level='1ST_CALL'));db.commit()
        content=b'name,call_level\nExisting Doctor,first call\n'
        first=preview_member_import(db,'first.csv',content,{})
        apply_member_import(db,'first.csv',content,{},first['preview_hash'],'internal-test')
        second=preview_member_import(db,'second.csv',content,{})
        assert first['preview_hash']!=second['preview_hash']
        result=apply_member_import(db,'second.csv',content,{},second['preview_hash'],'internal-test')
        assert not result['replayed']
        assert len(list(db.scalars(select(ImportBatch))))==2


def test_history_links_source_rows_and_exports_rejections(secured_client):
    import csv
    import io
    client,factory=secured_client
    token=login(client,'superadmin')['token'];headers={'Authorization':f'Bearer {token}'}
    with factory() as db:
        db.add_all([Person(canonical_name='Import Doctor',call_level='1ST_CALL'),Unit(code='UNIT_I',name='Unit I')]);db.commit()
    w=Workbook();w.active.title='=2+2';w.active.append(['Posting','Unit I']);w.active.append(['1st call','Import Doctor; Zebulon Zwart']);b=BytesIO();w.save(b)
    files={'file':('history.xlsx',b.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    preview=client.post('/api/v1/unit-management/import-preview?month=2026-10',headers=headers,files=files).json()
    result=client.post('/api/v1/unit-management/import-apply?month=2026-10',headers=headers,files=files,data={'preview_hash':preview['preview_hash']})
    assert result.status_code==200
    assert result.json()['created_rows']==1 and result.json()['skipped_rows']==1
    batch=result.json()['batch_id']
    history=client.get('/api/v1/imports/units/history?month=2026-10',headers=headers).json()
    assert len(history['items'])==1 and history['items'][0]['actor']=='superadmin'
    detail=client.get(f'/api/v1/imports/units/history/{batch}',headers=headers).json()
    assert len(detail['rows'])==2 and sum(bool(r['record_id']) for r in detail['rows'])==1
    assert client.get(f'/api/v1/imports/leave/history/{batch}',headers=headers).status_code==404
    report=client.get(f'/api/v1/imports/units/history/{batch}/rejected.csv',headers=headers)
    rejected=list(csv.DictReader(io.StringIO(report.text.lstrip('\ufeff'))))
    assert len(rejected)==1 and 'Zebulon Zwart' in rejected[0]['Original']
    assert rejected[0]['Sheet']=="'=2+2"
    hod=login(client,'hod')['token'];hod_headers={'Authorization':f'Bearer {hod}'}
    assert client.get('/api/v1/imports/members/history',headers=hod_headers).status_code==403
    assert client.get('/api/v1/imports/past/history',headers=hod_headers).status_code==403
    assert client.get('/api/v1/imports/units/history',headers=hod_headers).status_code==200
