"""Frozen operational imports with independent field normalization and row review."""

import csv
import hashlib
import io
import json
import re
from datetime import date, datetime, timedelta
from uuid import UUID

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    AdminMapping,
    DutyAssignment,
    DutySlot,
    ImportReviewBatch,
    LeaveRequest,
    MemberActivity,
    MemberCallLevelChange,
    Person,
    PersonAlias,
    PersonDesignation,
    PersonPosting,
    UnitCallMinimum,
    ReferenceValue,
    Unit,
)
from app.services.leave import month_bounds
from app.services.normalization import display_text, label_key, member_key, member_lookup, resolve_member, unit_key
from app.services.reference_catalog import resolve_value
from app.services.reference_registry import audit, lock_namespace

FIELDS = {
    "member_name": [
        "name",
        "person",
        "doctor",
        "doctor_name",
        "member",
        "staff_name",
        "canonical_name",
    ],
    "person_id": ["member_id", "person_id"],
    "unit_name": ["unit", "unit_name", "unit_label"],
    "unit_id": ["unit_id"],
    "posting_type": ["posting", "call_level", "call", "posting_type"],
    "starts_on": ["from", "start", "start_date", "starts_on", "date"],
    "ends_on": ["to", "end", "end_date", "ends_on"],
    "leave_type": ["type", "leave_type", "category"],
    "leave_slot": ["slot", "session", "leave_slot"],
    "status": ["status", "approval_status"],
    "notes": ["notes", "note", "remarks", "reason"],
}
HEADER_MAP = {label_key(h): key for key, headers in FIELDS.items() for h in [key, *headers]}


def fingerprint(db: Session) -> str:
    """A change anywhere in reference/availability context invalidates the preview."""
    state = []
    for model in [
        Person,
        PersonAlias,
        PersonDesignation,
        MemberCallLevelChange,
        Unit,
        UnitCallMinimum,
        AdminMapping,
        ReferenceValue,
        PersonPosting,
        LeaveRequest,
        DutySlot,
        DutyAssignment,
    ]:
        for row in db.execute(select(model.__table__).order_by(model.__table__.c.id)).mappings():
            state.append(
                (
                    model.__tablename__,
                    {
                        k: str(v) if not isinstance(v, (dict, list, type(None))) else v
                        for k, v in row.items()
                    },
                )
            )
    return hashlib.sha256(json.dumps(state, sort_keys=True, default=str).encode()).hexdigest()


def parse_date(raw: str, format: str) -> date:
    raw = display_text(raw)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?: 00:00:00)?", raw):
        return date.fromisoformat(raw[:10])
    parts = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", raw)
    if not parts:
        raise ValueError(f'Invalid date "{raw}"; use YYYY-MM-DD or choose the file date format')
    a, b, year = map(int, parts.groups())
    if format == "auto" and a <= 12 and b <= 12 and a != b:
        raise ValueError(f'Ambiguous date "{raw}"; select day/month/year or month/day/year')
    if format == "MDY" or (format == "auto" and b > 12):
        return date(year, a, b)
    return date(year, b, a)


def read_rows(
    db: Session, kind: str, filename: str, content: bytes, month: str, settings: dict
) -> tuple[list[str], list[dict]]:
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("Import file must be 5 MB or smaller")
    mapping = settings.get("mapping") or {}
    if not isinstance(mapping, dict):
        raise ValueError("Column mapping must be an object")
    tables = tabular_records(filename, content)
    if kind == "units" and not mapping and filename.lower().endswith(".csv"):
        first_line = content.decode("utf-8-sig").lstrip().splitlines()[0] if content.strip() else ""
        if first_line.upper().startswith(("MONTH:", "[UNIT:", "[PAIN]", "[SICU]", "[DRP]")):
            return legacy_rows(db, kind, filename, content, month)
    if tables:
        all_rows = []
        all_headers = []
        for sheet, records in tables:
            headers, rows = mapped_records(kind, filename, sheet, records, mapping)
            all_headers = list(dict.fromkeys([*all_headers, *headers]))
            all_rows.extend(rows)
        return all_headers, all_rows
    return legacy_rows(db, kind, filename, content, month)


def tabular_records(filename: str, content: bytes) -> list[tuple[str, list[list[str]]]]:
    """Retain date text until the operator selects the file convention."""
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("Import file must be 5 MB or smaller")
    if filename.lower().endswith(".csv"):
        try:
            records = list(
                csv.reader(io.StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
            )
        except (UnicodeDecodeError, csv.Error) as exc:
            raise ValueError("Upload a correctly quoted CSV UTF-8 file") from exc
        return [("CSV", records)]
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        try:
            tables = []
            for sheet in workbook:
                records = [
                    [str(v) if v is not None else "" for v in row]
                    for row in sheet.iter_rows(values_only=True)
                ]
                first = next((r for r in records if any(v.strip() for v in r)), [])
                if first and (any(label_key(h) in HEADER_MAP for h in first) or
                              len(first) >= 2 and all(v.strip() for v in first)):
                    tables.append((sheet.title, records))
            return tables
        finally:
            workbook.close()
    if filename.lower().endswith(".xls"):
        import pandas as pd
        frames = pd.read_excel(io.BytesIO(content), sheet_name=None, header=None)
        return [(name, [[str(v) if not pd.isna(v) else "" for v in row] for row in frame.values.tolist()])
                for name, frame in frames.items() if not frame.empty]
    return []


def mapped_records(kind: str, filename: str, sheet: str, records: list[list[str]], mapping: dict):
    indexed = [(n, r) for n, r in enumerate(records, 1) if any(v.strip() for v in r)]
    if not indexed:
        raise ValueError("The file is empty")
    headers = indexed[0][1]
    if len(set(label_key(h) for h in headers)) != len(headers):
        raise ValueError("Column headings must be unique")
    fields = [mapping.get(h, HEADER_MAP.get(label_key(h), "")) for h in headers]
    if "member_name" not in fields and "person_id" not in fields:
        raise ValueError("Map a column to member_name or person_id")
    used = [f for f in fields if f]
    if len(set(used)) != len(used) or any(f not in FIELDS for f in used):
        raise ValueError("Map each supported field at most once")
    rows = []
    for n, record in indexed[1:]:
        values = {f: v.strip() for f, v in zip(fields, record) if f}
        rows.append(
            {
                "key": f"{sheet}:{n}",
                "source": {"sheet": sheet, "row": n, "filename": filename},
                "original": dict(zip(headers, record)),
                "values": values,
                "parse_issues": ["Column count differs from headings"]
                if len(record) != len(headers)
                else [],
            }
        )
    return headers, rows


def legacy_rows(
    db: Session, kind: str, filename: str, content: bytes, month: str
) -> tuple[list[str], list[dict]]:
    if kind == "units":
        from app.services.unit_assignment_import import preview_unit_assignment_import

        preview = preview_unit_assignment_import(db, filename, content, month)
    else:
        from app.services.leave_import import preview_leave_import

        preview = preview_leave_import(db, filename, content, month)
    rows = []
    for index, row in enumerate(preview["rows"]):
        values = {
            "member_name": row.get("raw_person_name") or row.get("cleaned_person_name") or "",
            "unit_name": row.get("raw_unit_label") or row.get("unit_name") or "",
            "posting_type": row.get("raw_posting_label") or row.get("posting_type") or "",
            "starts_on": row.get("starts_on") or "",
            "ends_on": row.get("ends_on") or "",
            "leave_type": row.get("leave_type") or "ANNUAL_LEAVE",
            "leave_slot": row.get("leave_slot") or "FULL_DAY",
            "status": row.get("status") or "imported_pending_review",
            "notes": row.get("notes") or "",
        }
        rows.append(
            {
                "key": str(
                    row.get("row_key") or f"{row.get('sheet_name')}:{row.get('row_number')}:{index}"
                ),
                "source": {
                    "filename": filename,
                    "sheet": row.get("sheet_name"),
                    "row": row.get("row_number"),
                },
                "original": row,
                "values": values,
                "parse_issues": [],
            }
        )
    return [], rows


def resolve_unit(db: Session, raw: str, identity: str | None = None) -> dict:
    units = list(db.scalars(select(Unit)))
    by_id = {str(u.id): u for u in units}
    matches = {}
    for unit in units:
        labels = [unit.name, unit.code, f"{unit.campus or ''} {unit.name}"]
        for label in labels:
            key = unit_key(label)
            matches.setdefault(key, {})[str(unit.id)] = unit
    for mapping in db.scalars(
        select(AdminMapping).where(
            AdminMapping.mapping_type == "unit_label", AdminMapping.status == "reviewed"
        )
    ):
        unit = by_id.get(mapping.target_key) or next(
            (u for u in units if u.code == mapping.target_key), None
        )
        if unit:
            matches.setdefault(unit_key(mapping.source_label), {})[str(unit.id)] = unit
    candidates = list(matches.get(unit_key(raw), {}).values())
    unit = by_id.get(identity) if identity else candidates[0] if len(candidates) == 1 else None
    return {
        "unit_id": str(unit.id) if unit and unit.active_status == "active" else None,
        "name": unit.name if unit else None,
        "state": "matched"
        if unit and unit.active_status == "active"
        else "ambiguous"
        if len(candidates) > 1
        else "needs_review",
        "method": "unit_id" if identity and unit else "approved_unit_label" if unit else None,
        "candidates": [
            {"id": str(u.id), "name": u.name, "campus": u.campus, "status": u.active_status}
            for u in (candidates or units)
        ],
    }


def normalize_rows(
    db: Session, kind: str, month: str, rows: list[dict], settings: dict
) -> list[dict]:
    month_start, month_end = month_bounds(month)
    corrections = settings.get("corrections") or {}
    repeated = settings.get("repeated") or {}
    mode = settings.get("mode", "append")
    existing = list(db.scalars(select(PersonPosting if kind == "units" else LeaveRequest)))
    result = []
    identity_lookup = member_lookup(db, active_only=False)
    for row in rows:
        values = {**row["values"], **corrections.get(row["key"], {})}
        person_id = values.get("person_id") or repeated.get("members", {}).get(
            member_key(values.get("member_name", ""))
        )
        errors = list(row.get("parse_issues", []))
        try:
            person = resolve_member(
                db,
                str(values.get("member_name") or ""),
                UUID(str(person_id)) if person_id else None,
                as_of=month_start,
                lookup=identity_lookup,
            )
        except ValueError:
            person = {
                "person_id": None,
                "candidates": [],
                "state": "needs_review",
                "reason": "Invalid member ID",
            }
        if not person["person_id"]:
            errors.append("Select an active department member")
        after = {"person_id": person["person_id"], "notes": str(values.get("notes") or "")}
        try:
            start = parse_date(
                str(values.get("starts_on") or (month_start if kind == "units" else "")),
                settings.get("date_format", "auto"),
            )
            end = parse_date(
                str(values.get("ends_on") or (month_end if kind == "units" else start)),
                settings.get("date_format", "auto"),
            )
            if end < start:
                raise ValueError("End date precedes start date")
            if start > month_end or end < month_start:
                raise ValueError("Date range is outside the selected month")
            after.update(starts_on=str(start), ends_on=str(end))
        except ValueError as exc:
            errors.append(str(exc))
        unit = None
        try:
            if kind == "units":
                posting, reason = resolve_value(
                    db, "posting_type", str(values.get("posting_type") or "")
                )
                after["posting_type"] = posting
                if posting in {"PAIN", "SICU", "DRP"}:
                    after["unit_id"] = None
                else:
                    unit = resolve_unit(
                        db,
                        str(values.get("unit_name") or ""),
                        str(
                            values.get("unit_id")
                            or repeated.get("units", {}).get(
                                unit_key(str(values.get("unit_name") or ""))
                            )
                            or ""
                        )
                        or None,
                    )
                    after["unit_id"] = unit["unit_id"]
                    if not unit["unit_id"]:
                        errors.append("Select an active unit")
            else:
                for field, default in [
                    ("leave_type", "ANNUAL_LEAVE"),
                    ("leave_slot", "FULL_DAY"),
                    ("status", "imported_pending_review"),
                ]:
                    after[field] = resolve_value(
                        db,
                        "leave_status" if field == "status" else field,
                        str(values.get(field) or default),
                    )[0]
        except ValueError as exc:
            errors.append(str(exc))
        target = None
        if after.get("person_id") and after.get("starts_on"):
            targets = [
                r
                for r in existing
                if str(r.person_id) == after["person_id"]
                and str(r.starts_on) == after["starts_on"]
                and (kind == "units" or r.leave_slot == after.get("leave_slot"))
            ]
            if len(targets) > 1 and mode == "update_matching":
                errors.append("Multiple existing records match; edit the records individually")
            target = targets[0] if len(targets) == 1 else None
        before = record_values(target, kind) if target else None
        action = "update" if target and mode == "update_matching" else "create"
        if target and before == after:
            action = "unchanged"
        elif target and mode == "append":
            errors.append(
                "A record already exists for this member and date; use update matching or resolve the overlap"
            )
        if mode == "update_matching" and not target:
            errors.append("No unique existing record matches this row")
        warnings = []
        if after.get("person_id") and after.get("starts_on") and after.get("ends_on"):
            start, end = (
                date.fromisoformat(after["starts_on"]),
                date.fromisoformat(after["ends_on"]),
            )
            for record in existing:
                if record is target or str(record.person_id) != after["person_id"]:
                    continue
                if record.starts_on <= end and (record.ends_on is None or record.ends_on >= start):
                    if (
                        kind == "units"
                        and record.posting_type not in {"PAIN", "SICU", "DRP"}
                        and after.get("posting_type") not in {"PAIN", "SICU", "DRP"}
                    ):
                        if mode != "replace_month" or record.source != "unit_board":
                            errors.append("Overlapping primary posting exists")
                    if (
                        kind == "leave"
                        and record.status not in {"cancelled", "rejected"}
                        and after.get("status") not in {"cancelled", "rejected"}
                    ):
                        slot = after.get("leave_slot")
                        if slot == record.leave_slot or {"FULL_DAY", "CUSTOM"}.intersection(
                            {slot, record.leave_slot}
                        ):
                            if not (
                                mode == "replace_month"
                                and record.source in {"import", "normalized_import"}
                            ):
                                errors.append("Overlapping leave exists for the same session")
            if kind == "leave":
                conflicts = list(
                    db.scalars(
                        select(DutyAssignment)
                        .join(DutySlot)
                        .where(
                            DutyAssignment.person_id == UUID(after["person_id"]),
                            DutySlot.duty_date >= start,
                            DutySlot.duty_date <= end,
                        )
                    )
                )
                if conflicts:
                    warnings.append(f"{len(conflicts)} existing duties overlap this leave range")
        result.append(
            {
                **row,
                "values": values,
                "after": after,
                "before": before,
                "target_id": str(target.id) if target else None,
                "person_match": person,
                "unit_match": unit,
                "action": "needs_review" if errors else action,
                "issues": errors,
                "warnings": warnings,
                "remember_person_alias": bool(values.get("remember_person_alias")),
                "remember_unit_alias": bool(values.get("remember_unit_alias")),
                "normalizations": [
                    {
                        "field": field,
                        "original": row["values"].get(field, ""),
                        "proposed": proposed,
                        "reason": "Operator correction"
                        if field in corrections.get(row["key"], {})
                        else "Approved label or alias"
                        if field in {"posting_type", "leave_type", "leave_slot", "status"}
                        else "Selected month boundary"
                        if field in {"starts_on", "ends_on"} and not values.get(field)
                        else "Validated date or value",
                    }
                    for field, proposed in after.items()
                    if field not in {"person_id", "unit_id"}
                ],
            }
        )
    seen = {}
    excluded = set(settings.get("excluded", []))
    for row in result:
        if row["key"] in excluded:
            row["issues"].append("Excluded from this preview")
            row["action"] = "needs_review"
            continue
        after = row["after"]
        key = (
            after.get("person_id"),
            after.get("starts_on"),
            after.get("leave_slot") if kind == "leave" else after.get("posting_type"),
        )
        if key[0] and key in seen:
            for duplicate in [row, seen[key]]:
                duplicate["issues"].append("Repeated assignment or leave row in this file")
                duplicate["action"] = "needs_review"
        else:
            seen[key] = row
    for index, row in enumerate(result):
        if row["key"] in excluded:
            continue
        a = row["after"]
        if not all(a.get(f) for f in ["person_id", "starts_on", "ends_on"]):
            continue
        for other in result[index + 1 :]:
            if other["key"] in excluded:
                continue
            b = other["after"]
            if (
                a["person_id"] != b.get("person_id")
                or not b.get("starts_on")
                or not b.get("ends_on")
            ):
                continue
            if a["starts_on"] > b["ends_on"] or b["starts_on"] > a["ends_on"]:
                continue
            conflict = (
                kind == "units"
                and a.get("posting_type") not in {"PAIN", "SICU", "DRP"}
                and b.get("posting_type") not in {"PAIN", "SICU", "DRP"}
            )
            if kind == "leave":
                slots = {a.get("leave_slot"), b.get("leave_slot")}
                conflict = (
                    a.get("status") not in {"cancelled", "rejected"}
                    and b.get("status") not in {"cancelled", "rejected"}
                    and (len(slots) == 1 or "FULL_DAY" in slots or "CUSTOM" in slots)
                )
            if conflict:
                for item in [row, other]:
                    item["issues"].append("Overlapping rows for this member in this file")
                    item["action"] = "needs_review"
    return result


def record_values(record, kind: str) -> dict:
    fields = ["person_id", "starts_on", "ends_on", "notes"] + (
        ["unit_id", "posting_type"] if kind == "units" else ["leave_type", "leave_slot", "status"]
    )
    return {
        field: str(getattr(record, field))
        if isinstance(getattr(record, field), (UUID, date))
        else getattr(record, field) or ("" if field == "notes" else None)
        for field in fields
    }


def apply_batch(
    db: Session, batch: ImportReviewBatch, operator, excluded: list[str], acknowledge_warnings: bool
) -> dict:
    if batch.operator_id != operator.id:
        raise HTTPException(403, "This review belongs to another operator")
    if batch.result is not None:
        return batch.result
    if batch.created_at < datetime.utcnow() - timedelta(hours=2):
        raise HTTPException(409, "Review expired; preview the file again")
    lock_namespace(db, "operational-import")
    if batch.fingerprint != fingerprint(db):
        raise HTTPException(409, "Reference or operational records changed; preview again")
    selected = [r for r in batch.rows if r["key"] not in excluded]
    from app.services.auth import ADMIN_ROLES
    if operator.role not in ADMIN_ROLES and any(r["remember_person_alias"] or r["remember_unit_alias"] for r in selected):
        raise HTTPException(403, "Only an administrator can approve reusable identity aliases")
    if not selected:
        raise HTTPException(422, "Select at least one row")
    if any(r["issues"] for r in selected):
        raise HTTPException(422, "Resolve or exclude every invalid row")
    if any(r["warnings"] for r in selected) and not acknowledge_warnings:
        raise HTTPException(422, "Review and acknowledge the affected duties before applying leave")
    kind = batch.kind
    model = PersonPosting if kind == "units" else LeaveRequest
    result = {
        "created": 0,
        "updated": 0,
        "unchanged": 0,
        "removed": 0,
        "excluded": len(batch.rows) - len(selected),
        "batch_id": str(batch.id),
    }
    if batch.mode == "replace_month":
        start, end = month_bounds(batch.month)
        existing = list(
            db.scalars(
                select(model).where(
                    model.starts_on <= end, (model.ends_on.is_(None)) | (model.ends_on >= start)
                )
            )
        )
        for record in existing:
            if record.source not in (
                {"unit_board"} if kind == "units" else {"import", "normalized_import"}
            ):
                continue
            old = record_values(record, kind)
            audit(
                db,
                operator,
                "import_replacement",
                record.id,
                {"before": old, "replaced_month": batch.month, "batch_id": str(batch.id)},
            )
            if record.starts_on < start and (record.ends_on is None or record.ends_on > end):
                right = {
                    **old,
                    "starts_on": end + timedelta(days=1),
                    "ends_on": record.ends_on,
                    "person_id": record.person_id,
                }
                if kind == "units":
                    right["unit_id"] = record.unit_id
                db.add(model(**right, source=record.source))
                record.ends_on = start - timedelta(days=1)
            elif record.starts_on < start:
                record.ends_on = start - timedelta(days=1)
            elif record.ends_on is None or record.ends_on > end:
                record.starts_on = end + timedelta(days=1)
            else:
                db.delete(record)
            result["removed"] += 1
    for row in selected:
        retained_target = db.get(model, UUID(row["target_id"])) if row["target_id"] else None
        retained_manual = retained_target is not None and retained_target.source not in (
            {"unit_board"} if kind == "units" else {"import", "normalized_import"}
        )
        if row["action"] == "unchanged" and (batch.mode != "replace_month" or retained_manual):
            result["unchanged"] += 1
            continue
        record = (
            db.get(model, UUID(row["target_id"]))
            if row["target_id"] and batch.mode == "update_matching"
            else None
        )
        before = record_values(record, kind) if record else None
        if record is None:
            record = model(source="unit_board" if kind == "units" else "normalized_import")
            db.add(record)
        for key, value in row["after"].items():
            if key in {"person_id", "unit_id"}:
                value = UUID(value) if value else None
            if key in {"starts_on", "ends_on"}:
                value = date.fromisoformat(value) if value else None
            setattr(record, key, value)
        if kind == "leave":
            record.raw_person_name = str(row["values"].get("member_name") or "")
        db.flush()
        if row["remember_person_alias"]:
            from app.services.member_profile import ensure_unique_name

            alias = display_text(str(row["values"].get("member_name") or ""))
            person = db.get(Person, record.person_id)
            if alias and alias != person.canonical_name:
                ensure_unique_name(db, alias, person)
                if not any(member_key(a.alias) == member_key(alias) for a in person.aliases):
                    person.aliases.append(PersonAlias(alias=alias, source="reviewed_import"))
        if row["remember_unit_alias"] and kind == "units" and record.unit_id:
            # Alias writes stay inside this transaction; no service commit is invoked.
            alias = display_text(str(row["values"].get("unit_name") or ""))
            mappings = list(
                db.scalars(select(AdminMapping).where(AdminMapping.mapping_type == "unit_label"))
            )
            collisions = [m for m in mappings if unit_key(m.source_label) == unit_key(alias)]
            if any(m.target_key != str(record.unit_id) for m in collisions):
                raise HTTPException(409, "Unit alias belongs elsewhere")
            if alias and not collisions:
                db.add(
                    AdminMapping(
                        mapping_type="unit_label",
                        source_label=alias,
                        target_key=str(record.unit_id),
                        status="reviewed",
                        source="reviewed_import",
                    )
                )
        audit(
            db,
            operator,
            "posting_imported" if kind == "units" else "leave_imported",
            record.id,
            {
                "batch_id": str(batch.id),
                "source": row["source"],
                "raw": row["original"],
                "before": before,
                "after": row["after"],
                "person_match": row["person_match"],
                "normalizations": row.get("normalizations", []),
                "remember_person_alias": row["remember_person_alias"],
                "remember_unit_alias": row["remember_unit_alias"],
            },
        )
        db.add(
            MemberActivity(
                person_id=record.person_id,
                operator_id=operator.id,
                operator_name=operator.display_name or operator.username,
                action="posting_imported" if kind == "units" else "leave_imported",
                details={
                    "before": before,
                    "after": row["after"],
                    "source": row["source"],
                    "batch_id": str(batch.id),
                },
            )
        )
        result["updated" if before else "created"] += 1
    batch.result = result
    db.commit()
    return result
