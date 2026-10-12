"""Reviewed, non-destructive department member workbook imports."""

import hashlib
import csv
import json
import re
from datetime import date, datetime
from difflib import SequenceMatcher
from io import BytesIO, StringIO
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile
from uuid import UUID

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models import (
    ImportBatch,
    ImportSourceRecord,
    Person,
    PersonAlias,
    PersonDesignation,
    PersonCallChange,
)
from app.services.members import normalize_member_name
from app.domain.member_validation import normalize_member_display_name, normalize_member_call, normalize_member_status
from app.services.roster_reconciliation import normalize_designation
from app.services.imports import is_valid_person_name
from app.services.reviewed_import import acquire_import_lock, replayed_result, attach_result
from app.services.member_profiles import (
    call_level_on,
    local_today,
    record_call_change,
    sync_current_call_levels,
)

FIELDS = (
    "name",
    "designation",
    "call_level",
    "status",
    "aliases",
    "effective_from",
    "effective_to",
)
HEADER_ALIASES = {
    "name": {"name", "member name", "doctor name", "canonical name", "names"},
    "designation": {"designation", "position"},
    "call_level": {"call level", "call", "call levels"},
    "status": {"status", "active status"},
    "aliases": {"aliases", "alias"},
    "effective_from": {"effective from", "start date"},
    "effective_to": {"effective to", "end date"},
}


def cell_text(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return "" if value is None else str(value).strip()


def open_member_workbook(filename: str, content: bytes):
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("File must be no larger than 5 MB.")
    if filename.lower().endswith(".csv"):
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("Save the CSV using UTF-8 encoding.") from exc
        if "\x00" in decoded:
            raise ValueError("Invalid CSV text.")
        try:
            dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "CSV"
        try:
            for index, row in enumerate(csv.reader(StringIO(decoded), dialect, strict=True), 1):
                if index > 5100:
                    raise ValueError("Import at most 5000 data rows and 100 header rows.")
                if len(row) > 50:
                    raise ValueError("CSV files may contain at most 50 columns.")
                sheet.append(row)
            if not sheet.max_row or not decoded.strip():
                raise ValueError("CSV file is empty.")
        except (csv.Error, ValueError) as exc:
            workbook.close()
            raise ValueError(f"Cannot read CSV: {exc}") from exc
        return workbook
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Use a .csv or .xlsx file. Save older Excel files as .xlsx first.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Workbook expands beyond the 50 MB limit.")
        return load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, KeyError, OSError, ParseError) as exc:
        raise ValueError("Cannot read this Excel workbook.") from exc


def inspect_member_workbook(filename: str, content: bytes, header_row: int = 1):
    if not 1 <= header_row <= 100:
        raise ValueError("Header row must be between 1 and 100.")
    workbook = open_member_workbook(filename, content)
    try:
        sheets = []
        for sheet in workbook:
            values = next(
                sheet.iter_rows(
                    min_row=header_row,
                    max_row=header_row,
                    max_col=min(sheet.max_column or 50, 50),
                    values_only=True,
                ),
                (),
            )
            headers = [cell_text(value) for value in values]
            mapping = {}
            for field, aliases in HEADER_ALIASES.items():
                matches = [
                    i
                    for i, value in enumerate(headers)
                    if value.lower().replace("_", " ") in aliases
                ]
                if len(matches) == 1:
                    mapping[field] = matches[0]
            sheets.append({"name": sheet.title, "headers": headers, "mapping": mapping})
        return {"sheets": sheets, "header_row": header_row, "fields": FIELDS}
    finally:
        workbook.close()


def member_state(db: Session):
    sync_current_call_levels(db)
    people = list(
        db.scalars(
            select(Person).options(selectinload(Person.aliases), selectinload(Person.designations))
        )
    )
    snapshot = [
        {
            "id": str(p.id),
            "name": p.canonical_name,
            "status": p.active_status,
            "call": p.call_level,
            "archived": str(p.archived_at),
            "aliases": sorted(a.alias for a in p.aliases),
            "call_history": [
                (str(c.id), str(c.effective_from), c.to_call_level, c.source)
                for c in db.scalars(
                    select(PersonCallChange)
                    .where(PersonCallChange.person_id == p.id)
                    .order_by(PersonCallChange.id)
                )
            ],
            "designations": sorted(
                (d.designation, str(d.effective_from), str(d.effective_to)) for d in p.designations
            ),
        }
        for p in people
    ]
    return people, sorted(snapshot, key=lambda p: p["id"])


def preview_member_import(db: Session, filename: str, content: bytes, options: dict):
    header_row = options.get("header_row", 1)
    inspected = inspect_member_workbook(filename, content, header_row)
    sheet_name = options.get("sheet") or inspected["sheets"][0]["name"]
    sheet_info = next((s for s in inspected["sheets"] if s["name"] == sheet_name), None)
    if sheet_info is None:
        raise ValueError("Select an existing worksheet.")
    mapping = options.get("mapping", sheet_info["mapping"])
    if "name" not in mapping:
        raise ValueError("Map a column to member name.")
    if any(
        k not in FIELDS
        or not isinstance(v, int)
        or isinstance(v, bool)
        or v < 0
        or v >= len(sheet_info["headers"])
        for k, v in mapping.items()
    ):
        raise ValueError("Invalid column mapping.")
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("Each column can only be mapped once.")
    people, snapshot = member_state(db)
    by_id = {str(p.id): p for p in people}
    keys = {}
    for person in people:
        for name in [person.canonical_name] + [a.alias for a in person.aliases]:
            keys.setdefault(normalize_member_name(name), set()).add(str(person.id))
    resolutions = options.get("resolutions", {})
    rows, seen_names, seen_targets, seen_aliases = [], set(), set(), {}
    workbook = open_member_workbook(filename, content)
    try:
        sheet = workbook[sheet_name]
        for index, cells in enumerate(
            sheet.iter_rows(min_row=header_row + 1, max_col=50, values_only=True), header_row + 1
        ):
            if index - header_row > 5000:
                raise ValueError("Import at most 5000 rows at a time.")
            raw = {
                field: cell_text(cells[column]) if column < len(cells) else ""
                for field, column in mapping.items()
            }
            if not any(raw.values()):
                continue
            if resolutions.get(str(index)) == "skip":
                rows.append(
                    {
                        "row": index,
                        "raw": raw,
                        "normalized": {},
                        "action": "skip",
                        "target_id": None,
                        "target_name": None,
                        "changes": {},
                        "errors": [],
                        "candidates": [],
                    }
                )
                continue
            errors = []
            name = raw["name"]
            try:
                name = normalize_member_display_name(name)
            except ValueError as exc:
                errors.append(str(exc))
            if any(v.startswith("=") for v in raw.values()):
                errors.append("Formula cells are not supported; paste values first.")
            normalized = {
                "name": name,
                "designation": normalize_designation(raw.get("designation")),
                "status": raw.get("status", "").lower() or None,
                "call_level": None,
                "aliases": list(
                    dict.fromkeys(
                        a.strip() for a in re.split(r"[;\n]", raw.get("aliases", "")) if a.strip()
                    )
                ),
                "effective_from": raw.get("effective_from") or None,
                "effective_to": raw.get("effective_to") or None,
            }
            if raw["name"] != name and raw["name"] not in normalized["aliases"]:
                normalized["aliases"].append(raw["name"])
            try:
                if normalized["status"]:
                    normalized["status"] = normalize_member_status(normalized["status"])
                normalized["call_level"] = normalize_member_call(raw.get("call_level"))
            except ValueError as exc:
                errors.append(str(exc))
            for field in ("effective_from", "effective_to"):
                if normalized[field]:
                    try:
                        normalized[field] = date.fromisoformat(normalized[field]).isoformat()
                    except ValueError:
                        errors.append(f"{field} must be an Excel date or YYYY-MM-DD.")
            if normalized["designation"] and not normalized["effective_from"]:
                errors.append("Designation requires an effective-from date.")
            if not normalized["designation"] and (
                normalized["effective_from"] or normalized["effective_to"]
            ):
                errors.append("Effective dates require a designation.")
            if (
                normalized["effective_to"]
                and normalized["effective_from"]
                and normalized["effective_to"] < normalized["effective_from"]
            ):
                errors.append("Effective-to date precedes effective-from date.")
            if normalized["designation"] and len(normalized["designation"]) > 100:
                errors.append("Designation exceeds 100 characters.")
            if any(len(a) > 255 or not is_valid_person_name(a) for a in normalized["aliases"]):
                errors.append("Invalid alias or alias exceeds 255 characters.")
            key = normalize_member_name(name)
            if key in seen_names or key in seen_aliases:
                errors.append("Duplicate name in this worksheet; consolidate these rows.")
            seen_names.add(key)
            matches = keys.get(key, set())
            candidates = set(matches)
            if not matches:
                candidates = {
                    str(p.id)
                    for p in people
                    if SequenceMatcher(None, key, normalize_member_name(p.canonical_name)).ratio()
                    >= 0.82
                }
            resolution = resolutions.get(str(index), "")
            target = next(iter(matches)) if len(matches) == 1 else None
            if resolution and resolution not in {"new", "skip"}:
                if resolution not in by_id:
                    errors.append("Selected member no longer exists.")
                else:
                    target = resolution
            if resolution == "new":
                if matches:
                    errors.append("Name already matches an existing member.")
                target = None
            if target and target in seen_targets:
                errors.append("Multiple rows target the same member; consolidate these rows.")
            if target:
                seen_targets.add(target)
                if name != by_id[target].canonical_name and name not in normalized["aliases"]:
                    normalized["aliases"].append(name)
            for alias in normalized["aliases"]:
                alias_key = normalize_member_name(alias)
                owners = keys.get(alias_key, set())
                if owners - ({target} if target else set()):
                    errors.append(f"Alias belongs to another member: {alias}")
                if alias_key in seen_aliases and seen_aliases[alias_key] != index:
                    errors.append(f"Alias repeated across workbook rows: {alias}")
                if alias_key in seen_names and alias_key != key:
                    errors.append(f"Alias conflicts with another workbook member: {alias}")
                seen_aliases[alias_key] = index
            changes = {}
            person = by_id.get(target)
            if person:
                for field, attr in (("status", "active_status"), ("call_level", "call_level")):
                    if normalized[field] and normalized[field] != getattr(person, attr):
                        changes[field] = {"from": getattr(person, attr), "to": normalized[field]}
                new_aliases = [
                    a
                    for a in normalized["aliases"]
                    if a != person.canonical_name
                    and a not in {alias.alias for alias in person.aliases}
                ]
                if new_aliases:
                    changes["aliases"] = new_aliases
                if normalized["designation"]:
                    designation = next(
                        (
                            d
                            for d in person.designations
                            if d.designation == normalized["designation"]
                            and str(d.effective_from) == normalized["effective_from"]
                        ),
                        None,
                    )
                    if designation is None:
                        changes["designation"] = normalized["designation"]
                    elif (
                        normalized["effective_to"]
                        and str(designation.effective_to) != normalized["effective_to"]
                    ):
                        errors.append(
                            "Existing designation has a different end date; review manually."
                        )
            action = "update" if person and changes else "unchanged" if person else "create"
            if errors or ((len(matches) > 1 or (not matches and candidates)) and not resolution):
                action = "needs_review"
            if resolution == "skip":
                action = "skip"
            rows.append(
                {
                    "row": index,
                    "raw": raw,
                    "normalized": normalized,
                    "action": action,
                    "target_id": target,
                    "target_name": person.canonical_name if person else None,
                    "changes": changes,
                    "errors": errors,
                    "candidates": [
                        {"id": str(p.id), "name": p.canonical_name}
                        for p in people
                        if str(p.id) in candidates
                    ],
                }
            )
    finally:
        workbook.close()
    # Include the entire member state so a concurrent edit invalidates the review.
    digest = hashlib.sha256(
        json.dumps([rows, snapshot, sheet_name, header_row, filename, hashlib.sha256(content).hexdigest(),
                    {key:value for key,value in options.items() if key != "confirm_call_changes"}], sort_keys=True).encode()
    ).hexdigest()
    counts = {
        action: sum(r["action"] == action for r in rows)
        for action in ("create", "update", "unchanged", "needs_review", "skip")
    }
    return {"rows": rows, "counts": counts, "preview_hash": digest, "sheet": sheet_name}


def apply_member_import(
    db: Session, filename: str, content: bytes, options: dict, preview_hash: str, actor_id: str
):
    acquire_import_lock(db)
    replay = replayed_result(db, "department_members", preview_hash, actor_id, filename, content, options)
    if replay is not None:
        return replay
    preview = preview_member_import(db, filename, content, options)
    if preview["preview_hash"] != preview_hash:
        raise ValueError("File, selections, or members changed. Preview again before importing.")
    if preview["counts"]["needs_review"]:
        raise ValueError("Resolve or skip all rows needing review before importing.")
    if any("call_level" in row["changes"] for row in preview["rows"] if row["action"] == "update"):
        if options.get("confirm_call_changes") is not True:
            raise ValueError("Confirm the reviewed call-level changes before importing.")
    batch = ImportBatch(
        source_filename=filename[:255],
        import_kind="department_members",
        status="completed",
        completed_at=datetime.utcnow(),
        source_metadata={
            "actor_id": actor_id,
            "counts": preview["counts"],
            "preview_hash": preview_hash,
            "file_sha256": hashlib.sha256(content).hexdigest(),
            "sheet": preview["sheet"],
            "call_changes_confirmed": options.get("confirm_call_changes") is True,
        },
    )
    db.add(batch)
    try:
        for row in preview["rows"]:
            data = row["normalized"]
            person = db.get(Person, UUID(row["target_id"])) if row["target_id"] else None
            if row["action"] == "create":
                person = Person(
                    canonical_name=data["name"],
                    active_status=data["status"] or "active",
                    call_level=data["call_level"],
                )
                db.add(person)
                db.flush()
            if row["action"] in {"create", "update"}:
                if data["status"]:
                    person.active_status = data["status"]
                person.archived_at = (
                    (person.archived_at or datetime.utcnow())
                    if person.active_status == "archived"
                    else None
                )
                if data["call_level"]:
                    if data["call_level"] != call_level_on(person):
                        record_call_change(
                            db,
                            person,
                            data["call_level"],
                            local_today(),
                            f"Reviewed upload: {filename}, row {row['row']}",
                            "member_excel_import",
                            actor_id=UUID(actor_id)
                            if re.fullmatch(r"[0-9a-fA-F-]{36}", actor_id)
                            else None,
                        )
                for alias in data["aliases"]:
                    if alias != person.canonical_name and alias not in {
                        a.alias for a in person.aliases
                    }:
                        person.aliases.append(
                            PersonAlias(alias=alias, source="member_excel_import")
                        )
                if data["designation"] and not any(
                    d.designation == data["designation"]
                    and str(d.effective_from) == data["effective_from"]
                    for d in person.designations
                ):
                    person.designations.append(
                        PersonDesignation(
                            designation=data["designation"],
                            effective_from=date.fromisoformat(data["effective_from"]),
                            effective_to=date.fromisoformat(data["effective_to"])
                            if data["effective_to"]
                            else None,
                            source="member_excel_import",
                        )
                    )
            db.add(
                ImportSourceRecord(
                    batch=batch,
                    sheet_name=preview["sheet"],
                    row_index=row["row"],
                    record_type="department_member",
                    raw_value=json.dumps(row["raw"]),
                    cleaned_value=json.dumps(data),
                    normalized_table="persons" if person else None,
                    normalized_id=person.id if person else None,
                    rule_key="member_excel_import_v1",
                    notes=json.dumps(
                        {
                            "action": row["action"],
                            "changes": row["changes"],
                            "errors": row["errors"],
                        }
                    ),
                )
            )
        db.flush()
        result = {"batch_id": str(batch.id), "counts": preview["counts"], "replayed": False}
        attach_result(batch, result, filename, content, options)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return result
